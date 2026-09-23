"""Update the existing submission DOCX without rebuilding its page layout.

Run only after auditing evidence/paper-corrected-issues5-12-externalrss-10ms-20260923. The current DOCX
is the source, so submission-specific layout edits made after the original
paper builder are preserved.
"""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from build_final_research_paper import (
    make_calibration_chart,
    make_diff_chart,
    make_operational_chart,
    make_sensitivity_chart,
)


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "evidence" / "paper-corrected-issues5-12-externalrss-10ms-20260923"
SENSITIVITY = ROOT / "evidence" / "trie-sensitivity-issues8-20260923" / "summary.csv"
PAPER = ROOT / "paper" / "Revon_Final_Research_Paper.docx"
COVER = ROOT / "output" / "docs" / "Revon_Cover_Letter_Computing.docx"
TEMP = ROOT / "tmp" / "paper_first_four_issues"
SCENARIOS = ("small-sparse", "medium-sparse", "medium-dense", "large-sparse", "large-hot")
LABELS = ("Small sparse", "Medium sparse", "Medium dense", "Large sparse", "Large hot")


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def number(row: dict[str, str], metric: str) -> float:
    return float(row[f"{metric}_median"])


def main() -> None:
    manifest = json.loads((EVIDENCE / "manifest.json").read_text(encoding="utf-8"))
    audit = json.loads((EVIDENCE / "evidence_audit.json").read_text(encoding="utf-8"))
    raw = read_csv(EVIDENCE / "raw_results.csv")
    summary = read_csv(EVIDENCE / "summary.csv")
    assert manifest["schema_version"] == 3 and manifest["profile"] == "paper"
    assert audit["passed"] and not audit["issues"]
    assert audit["execution_counts"] == {
        "warmups": 74,
        "measured": 259,
        "primary_evaluation_measured": 175,
        "dolt_evaluation_measured": 35,
    }
    assert len(raw) == 333 and len(summary) == 37
    assert all(row["status"] == "ok" and row["correctness"] == "True" for row in raw)
    assert all(row["changed_keys"] for row in raw if row["model"] == "Dolt")
    assert manifest["hybrid_threshold_operations"] == 4096, "Update the calibration chart for a changed threshold"
    evaluation = {(row["scenario"], row["model"]): row for row in summary if row["phase"] == "evaluation"}
    calibration = [row for row in summary if row["phase"] == "calibration"]

    def row(scenario: str, model: str) -> dict[str, str]:
        return evaluation[(scenario, model)]

    def speedup(scenario: str, metric: str, baseline: str = "Dolt") -> float:
        return number(row(scenario, baseline), metric) / number(row(scenario, "Revon-H"), metric)

    def span(values: list[float]) -> str:
        return f"{min(values):.1f} to {max(values):.1f}"

    def performance(values: list[float], baseline: str = "Dolt") -> str:
        if min(values) > 1:
            return f"{span(values)} times as fast as {baseline}"
        if max(values) < 1:
            inverse = [1 / value for value in values]
            return f"{span(inverse)} times slower than {baseline}"
        return (
            f"{baseline}/Revon-H median-latency ratios from {span(values)} "
            "(ratios above one favor Revon-H)"
        )

    diff = [speedup(s, "diff_ms") for s in SCENARIOS]
    commit = [speedup(s, "incremental_commit_ms") for s in SCENARIOS]
    checkout = [speedup(s, "checkout_ms") for s in SCENARIOS]
    storage = [number(row(s, "Revon-H"), "storage_bytes") / number(row(s, "Dolt"), "storage_bytes") for s in SCENARIOS]
    ablation = [number(row(s, "Revon-M (forced Merkle)"), "diff_ms") / number(row(s, "Revon-H"), "diff_ms") for s in SCENARIOS]
    dense_gap = abs(number(row("medium-dense", "Revon-H"), "diff_ms") / number(row("medium-dense", "Revon-M (forced Merkle)"), "diff_ms") - 1) * 100
    large_import = [number(row(s, "Revon-H"), "initial_import_ms") / number(row(s, "Dolt"), "initial_import_ms") for s in SCENARIOS[-2:]]

    memory_models = ("Snapshot", "Log-only", "Revon-M (forced Merkle)", "Revon-H", "Dolt")
    memory_ranges = []
    for model in memory_models:
        values = [number(row(scenario, model), "process_tree_peak_rss_bytes") / (1024 * 1024) for scenario in SCENARIOS]
        memory_ranges.append(f"{model}: {min(values):.1f}-{max(values):.1f} MiB")
    memory_result = "; ".join(memory_ranges)
    sensitivity = read_csv(SENSITIVITY)
    sensitivity_by_config = {row["config"]: row for row in sensitivity}
    assert len(sensitivity) == 5 and all(row["all_correct"] == "True" for row in sensitivity)
    b8d4 = sensitivity_by_config["b8-d4"]
    b8d5 = sensitivity_by_config["b8-d5"]
    fastest = min(sensitivity, key=lambda row: float(row["diff_ms_median"]))
    storage_ratio_b8d5 = float(b8d5["storage_bytes_median"]) / float(b8d4["storage_bytes_median"])
    sensitivity_text = (
        "The corrected fixed-trie sensitivity run contains 45 successful executions: 10 warm-ups and 35 measured runs. "
        f"{fastest['config']} had the lowest median diff ({float(fastest['diff_ms_median']):.2f} ms) but used "
        f"{float(fastest['storage_bytes_median']) / (1024 * 1024):.2f} MiB; b8-d4 measured "
        f"{float(b8d4['diff_ms_median']):.2f} ms and used the least storage ({float(b8d4['storage_bytes_median']) / (1024 * 1024):.2f} MiB). "
        f"The over-deep b8-d5 design measured {float(b8d5['diff_ms_median']):.2f} ms for diff and used "
        f"{(storage_ratio_b8d5 - 1) * 100:.1f}% more storage than b8-d4. Thus b8-d4 is a storage-conscious balanced default, not a universal optimum."
    )

    doc = Document(PAPER)

    def replace(prefix: str | tuple[str, ...], text: str) -> None:
        prefixes = (prefix,) if isinstance(prefix, str) else prefix
        matches = [p for p in doc.paragraphs if p.text.startswith(prefixes)]
        assert len(matches) == 1, (prefix, len(matches))
        paragraph = matches[0]
        assert paragraph.runs, prefix
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""

    replace(
        "Versioning structured datasets requires",
        "Versioning structured datasets requires durable history, efficient sparse updates, historical reconstruction, and comparison without copying every record. Revon is a Python prototype built around an immutable fixed-depth Merkle hash trie. A commit rewrites affected leaf buckets and their ancestor paths while reusing unchanged subtrees by content hash. Revon-H adds addressed changesets and selects log aggregation or hash-pruned tree differencing with a separately calibrated threshold. We evaluate Snapshot, Log-only, forced-Merkle Revon-M, Revon-H, and Dolt 2.3.1 on deterministic workloads up to 100,000 rows. The corrected run contains 333 executions: 74 warm-ups and 259 measured runs, including 175 primary evaluation runs. In this in-process Revon versus Dolt CLI workflow, "
        f"Dolt/Revon-H median-latency ratios were {span(diff)} for diff and {span(commit)} for incremental commit (ratios above one favor Revon-H). "
        f"Revon-H/Dolt median repository-footprint ratios were {span(storage)} (values above one mean larger Revon-H repository footprints). Against forced Merkle, the hybrid log path was "
        f"{performance([ablation[i] for i in (0, 1, 3, 4)], 'forced-Merkle Revon-M')} on sparse or hot diffs. "
        "These workflow measurements support a workload-dependent hybrid design, not general superiority over a production SQL database."
    )
    replace(
        "A five-variant benchmark with shared",
        "A five-variant benchmark with shared deterministic workloads, common canonical diff and checkout outputs, process-tree RSS sampling, raw evidence, and an internal evidence-integrity audit.",
    )
    comparator = next(p for p in doc.paragraphs if p.text.startswith("Dolt is the production comparator because"))
    assert len(comparator._p.xpath('.//w:hyperlink')) == 2 and len(comparator.runs) == 3
    comparator.runs[2].text = ". Revon is not presented as a feature replacement for Dolt. The comparison measures a smaller in-process Python prototype against a Dolt CLI workflow and does not isolate their tree algorithms."
    replace(
        "Five implementations receive identical logical states",
        "Five implementations receive identical logical states and mutation batches: a complete Snapshot per version, an initial state plus operation-log baseline, forced-Merkle Revon-M, adaptive Revon-H, and Dolt 2.3.1 using a keyed SQL table and native JSON diff. Revon-M isolates the selector.",
    )
    replace(
        ("Initial import measures creation of the first", "Initial import measures the first durable state"),
        "Initial import measures the first durable state; incremental commit is the within-trial median across ten commits; diff and checkout both materialize common outputs inside their timed regions. Diff is sorted changed keys with old/new SHA-256 value hashes encoded as canonical JSON. Checkout is the complete sorted key/value state encoded as canonical UTF-8 JSON. Dolt JSON parsing and normalization are included. Correctness compares the full state and hash contract with the workload oracle. Dolt commit timing includes SQL mutation, add, and commit; commit hashes are retained for historical checkout and no version tags are created. An external parent process samples process-tree RSS every 10 ms for each isolated trial; sampling is outside timed operations. Work-examined units differ by system and are not normalized.",
    )
    replace(
        ("Before measurement, each configuration runs twice", "Each configuration has two warm-ups"),
        f"Each configuration has two warm-ups and seven measured runs. Across calibration and evaluation, this gives 74 warm-ups and 259 measured runs. The primary five-system evaluation contains 175 measured runs, including 35 Dolt runs. Each system/trial runs in a clean worker process. One deterministic workload is reused per scenario; the internal evidence-integrity audit regenerates it from seed {manifest['base_seed']} and checks its digest. Summaries exclude warm-ups. Run {manifest['run_id']} was measured on {manifest['platform']} with Python {manifest['python_version']}. The audit recomputes the trial matrix and summary statistics; it does not independently reproduce timings.",
    )
    replace(
        ("The boundary is a frozen policy", "The 4,096-operation boundary is"),
        f"The 4,096-operation boundary is a frozen policy for this machine, not a universal crossover. All 333 executions completed successfully and passed the harness's state and diff checks. The internal audit retained all {audit['outlier_candidates']} Tukey outlier candidates; no observation was manually deleted.",
    )
    replace(
        ("Revon-H selected log for the four sparse", "Revon-H selected log for four sparse"),
        f"Revon-H selected log for four sparse or hot workloads and Merkle for medium dense. Its median diff was {performance([diff[2]])} in medium dense and {performance([diff[i] for i in (0, 1, 3, 4)])} in the others. The comparison uses Dolt median divided by Revon-H median for each scenario; values are ratios of separate medians without confidence intervals. Both systems materialize the same sorted key/hash JSON result. Dolt's CLI process and SQL costs remain inside this workflow comparison.",
    )
    replace(
        ("For log-selected workloads, Revon-H reduced", "For log-selected workloads, Revon-H's median diff"),
        f"For log-selected workloads, Revon-H's median diff was {performance([ablation[0]], 'forced-Merkle Revon-M')} in small sparse, {performance([ablation[1]], 'forced-Merkle Revon-M')} in medium sparse, {performance([ablation[3]], 'forced-Merkle Revon-M')} in large sparse, and {performance([ablation[4]], 'forced-Merkle Revon-M')} in large hot. In medium dense, both used Merkle; their median difference was {dense_gap:.1f}%. This ablation supports the selector's intended mechanism on the tested workloads.",
    )
    replace(
        ("Revon-H incremental commit medians were", "Revon-H's median incremental commit was", "The Dolt/Revon-H median-latency ratio ranged from"),
        f"The Dolt/Revon-H median-latency ratio ranged from {span(commit)} for incremental commit and {span(checkout)} for historical checkout. In the two 100,000-row scenarios, Revon-H initial-import latency was {large_import[0]:.2f} and {large_import[1]:.2f} times Dolt's. These are workflow ratios, not isolated tree-algorithm speedups.",
    )
    replace(
        "Revon-H paid a consistent storage cost",
        "Revon-H paid a consistent storage cost. Its median repository-footprint ratios against Dolt were "
        + ", ".join(f"{value:.2f}" for value in storage[:-1])
        + f", and {storage[-1]:.2f} from small sparse through large hot, respectively; values above 1 mean a larger Revon-H repository. "
        + f"The two 100,000-row Revon repositories occupied about {sum(number(row(s, 'Revon-H'), 'storage_bytes') for s in SCENARIOS[-2:]) / (2 * 1024 * 1024):.1f} MiB on average. This is total workflow footprint, including each system's encoding and metadata.",
    )
    replace(
        ("All 45 sensitivity trials were correct", "The sensitivity bundle contains 45 successful executions", "The corrected fixed-trie sensitivity run contains"),
        sensitivity_text,
    )
    replace(
        ("The results answer RQ1 with a trade-off", "RQ1 shows a workflow trade-off", "RQ1 shows scenario-dependent workflow latency"),
        "RQ1 shows scenario-dependent workflow latency and storage trade-offs. Revon-H's measured median ratios vary by operation and dataset, and its two 100,000-row imports have the ratios reported above. Snapshot and Log-only remain useful representation baselines. These results compare the declared interfaces and do not isolate storage-tree algorithms.",
    )
    replace(
        "Outlier policy: all",
        f"Outlier policy: all {audit['outlier_candidates']} Tukey candidates were kept. Medians and interquartile ranges limit their influence, but seven measured trials do not demonstrate normality or provide ratio confidence intervals.",
    )
    replace(
        "Figure 6: Fixed-trie sensitivity",
        "Figure 6: Fixed-trie sensitivity; dots show median diff latency, and the axis extends above the largest measured p75.",
    )
    replace(
        ("Memory comparability:", "Memory measurement:", "Memory result:"),
        "Memory result: Across the five primary-evaluation scenarios, median sampled whole-worker process-tree RSS ranged from "
        + memory_result
        + ". These readings include common workload/oracle state and adapter processes, so they are not product-only memory. The parent sampled every 10 ms outside timed operations; peaks shorter than 10 ms may be missed, and the monitor can create small system-level scheduling pressure.",
    )
    replace(
        ("Revon demonstrates a complete content-addressed", "Revon demonstrates a content-addressed"),
        "Revon demonstrates a content-addressed versioning path for structured key/value data: canonical immutable objects, incremental Merkle trie updates, atomic SQLite persistence, historical checkout, and adaptive differencing. In 175 primary measured evaluation runs across five systems and five scenarios, the paper reports scenario-level latency ratios, import costs, storage, and sampled process-tree RSS under the declared workflows. The within-system ablation provides more direct evidence about hybrid log/Merkle selection than cross-system ratios. These measurements are a reproducible prototype comparison, not a claim of production-system replacement.",
    )
    replace(
        "To move beyond a linear prototype",
        "To move beyond a linear prototype, Revon still needs named branches, merge commits, conflict reporting, schema-aware records, concurrent readers and writers, and remote synchronization. Evaluation should add public structured datasets, one-million-row and longer-history workloads, Linux replication, and unified external CPU and I/O telemetry. Current RSS figures are workflow-level samples and do not cover throughput, concurrency, or range-query behavior.",
    )
    replace(
        "The repository contains the benchmark harness",
        "The repository contains the benchmark harness, evidence audit, trie-sensitivity harness, and manuscript-formatting and validation tools. Source code, artifacts, and reproduction instructions are available at https://github.com/atharvasheersh/Revon. The corrected paper-profile evidence bundle is evidence/paper-corrected-issues5-12-externalrss-10ms-20260923; the trie-sensitivity bundle is evidence/trie-sensitivity-issues8-20260923. Tables and figures are produced from the audited summary CSV files and verified against raw results and manifests; no table or graph uses an earlier validation export.",
    )

    table = doc.tables[2]
    assert len(table.rows) == 6 and table.cell(0, 0).text == "Scenario"
    def set_cell(cell, value: str) -> None:
        paragraph = cell.paragraphs[0]
        assert paragraph.runs
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""

    set_cell(table.cell(0, 4), "Speedup")
    for index, (scenario, label) in enumerate(zip(SCENARIOS, LABELS), start=1):
        h, d = row(scenario, "Revon-H"), row(scenario, "Dolt")
        values = (label, h["strategy_selected"], f"{number(h, 'diff_ms'):.3f}", f"{number(d, 'diff_ms'):.2f}", f"{speedup(scenario, 'diff_ms'):.1f}x")
        for column, value in enumerate(values):
            set_cell(table.cell(index, column), value)

    TEMP.mkdir(parents=True, exist_ok=True)
    assets = {
        "Figure 3:": ("calibration.png", lambda path: make_calibration_chart(path, calibration)),
        "Figure 4:": ("diff.png", lambda path: make_diff_chart(path, evaluation)),
        "Figure 5:": ("operational.png", lambda path: make_operational_chart(path, evaluation)),
        "Figure 6:": ("sensitivity.png", lambda path: make_sensitivity_chart(path, sensitivity)),
    }
    for caption, (filename, render) in assets.items():
        caption_index = next(i for i, paragraph in enumerate(doc.paragraphs) if paragraph.text.startswith(caption))
        image_paragraph = next(paragraph for paragraph in reversed(doc.paragraphs[:caption_index]) if paragraph._p.xpath('.//a:blip'))
        blips = image_paragraph._p.xpath('.//a:blip')
        assert len(blips) == 1
        rid = blips[0].get(qn("r:embed"))
        image_path = TEMP / filename
        render(image_path)
        doc.part.related_parts[rid]._blob = image_path.read_bytes()

    backup = TEMP / "before_revisions.docx"
    shutil.copy2(PAPER, backup)
    output = TEMP / PAPER.name
    doc.save(output)
    output.replace(PAPER)
    print(PAPER)

    if COVER.exists():
        cover = Document(COVER)
        paragraphs = [p for p in cover.paragraphs if p.text.startswith("We evaluate Revon against snapshot")]
        assert len(paragraphs) == 1 and paragraphs[0].runs
        paragraphs[0].runs[0].text = (
            "We evaluate Revon against snapshot and log-only baselines, a forced-Merkle ablation, "
            "and Dolt 2.3.1 on deterministic workloads of up to 100,000 rows. The corrected run "
            "contains 333 executions: 74 warm-ups and 259 measured runs, including 175 in the "
            "primary five-system evaluation. Dolt's structured diff keys and values are checked "
            "against the workload oracle. Under the tested in-process Revon versus Dolt CLI "
            "workflow, Revon-H had lower median diff, commit, and checkout latency in the "
            "reported scenarios, with higher repository footprint; in the two 100,000-row "
            "scenarios, Revon-H initial import was 0.79x and 0.56x Dolt's median time. "
            "The manuscript reports these workflow trade-offs without claiming universal "
            "superiority over a production SQL system."
        )
        for run in paragraphs[0].runs[1:]:
            run.text = ""
        shutil.copy2(COVER, TEMP / "before_revisions_cover_letter.docx")
        cover_output = TEMP / COVER.name
        cover.save(cover_output)
        cover_output.replace(COVER)
        print(COVER)


if __name__ == "__main__":
    main()
