"""Record publication expectations from audited inputs and rendered artifacts."""
from pathlib import Path
from hashlib import sha256
import json
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
def sha(data):return sha256(data).hexdigest()

def main():
    inputs=json.loads((ROOT/'output/review_revision/publication_inputs.json').read_text())
    doc=Document(ROOT/'paper/Revon Research Paper Springer Nature.docx')
    inputs['tables_cells']=[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
    inputs['figure_sha256']={str(i):sha(doc.part.related_parts[s._inline.xpath('.//a:blip')[0].get(qn('r:embed'))].blob) for i,s in enumerate(doc.inline_shapes,1)}
    inputs['reference_paragraphs']=[p.text for p in doc.paragraphs if __import__('re').match(r'^\[\d+\]',p.text)]
    inputs['body_paragraphs']=[p.text for p in doc.paragraphs if p.text.strip()]
    inputs['abstract']=next(p.text for p in doc.paragraphs if p.text.startswith('Revon is a Python '))
    inputs['archive_name']='Revon_Review_Revision_Reproducibility_20261016.zip'
    inputs['campaign']='evidence/review-revision-20261016/campaign'
    inputs['analysis_sha256']=sha((ROOT/inputs['campaign']/'analysis.json').read_bytes())
    inputs['analysis_tables_rule']='numeric tables derive from audited analysis; historical cells retain independent baseline checks with citation renumbering and one exact source-description wording replacement'
    path=ROOT/'paper/Revon_Review_Publication_Model_20261016.json'
    path.write_text(json.dumps(inputs,indent=2),encoding='utf8')
    print(path)

if __name__=='__main__':main()
