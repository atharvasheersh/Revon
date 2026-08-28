from __future__ import annotations

import csv
import json
import math
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
FINAL_DIR = ROOT / "evidence" / "paper-final-20260824"
SENSITIVITY_DIR = ROOT / "evidence" / "trie-sensitivity-20260824"
OUT_DOCX = ROOT / "paper" / "Chronos_Final_Research_Paper.docx"
ASSET_DIR = ROOT / "tmp" / "final_paper_assets"

INK = "000000"
MUTED = "52606D"
GRID = "000000"
LIGHT = "FFFFFF"
BLUE = "2F6B9A"
BLUE_LIGHT = "EAF3F8"
GREEN = "4F8A5B"
AMBER = "D98B2B"
RED = "C94F4F"
WHITE = "FFFFFF"
HYPERLINK_BLUE = "0000FF"
COLORS = {
    "Snapshot": "4C78A8",
    "Log-only": "F58518",
    "Chronos-M (forced Merkle)": "E45756",
    "Chronos-H": "54A24B",
    "Dolt": "B279A2",
}
SENSITIVITY_COLORS = {
    "b4-d6": "4C78A8",
    "b8-d3": "F58518",
    "b8-d4": "54A24B",
    "b8-d5": "E45756",
    "b16-d3": "B279A2",
}


def column_widths_from_weights(weights, total_width_dxa):
    if not weights or any(weight <= 0 for weight in weights):
        raise ValueError("table weights must be positive")
    total_weight = float(sum(weights))
    widths = [
        int(round(total_width_dxa * weight / total_weight)) for weight in weights
    ]
    widths[-1] += total_width_dxa - sum(widths)
    return widths


def ensure_ooxml_child(parent, tag):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_ooxml_width(parent, tag, width_dxa):
    node = ensure_ooxml_child(parent, tag)
    node.set(qn("w:type"), "dxa")
    node.set(qn("w:w"), str(int(width_dxa)))


def apply_table_geometry(table, column_widths_dxa, *, table_width_dxa, indent_dxa):
    widths = [int(width) for width in column_widths_dxa]
    if sum(widths) != table_width_dxa:
        raise ValueError("column widths must equal the table width")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_properties = table._tbl.tblPr
    set_ooxml_width(table_properties, "w:tblW", table_width_dxa)
    table_indent = ensure_ooxml_child(table_properties, "w:tblInd")
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(indent_dxa))
    layout = ensure_ooxml_child(table_properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)
    for row in table.rows:
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            set_ooxml_width(cell._tc.get_or_add_tcPr(), "w:tcW", width)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def load_evidence():
    manifest = json.loads((FINAL_DIR / "manifest.json").read_text(encoding="utf-8"))
    audit = json.loads((FINAL_DIR / "evidence_audit.json").read_text(encoding="utf-8"))
    raw = read_csv(FINAL_DIR / "raw_results.csv")
    summary = read_csv(FINAL_DIR / "summary.csv")
    sensitivity_manifest = json.loads(
        (SENSITIVITY_DIR / "manifest.json").read_text(encoding="utf-8")
    )
    sensitivity = read_csv(SENSITIVITY_DIR / "summary.csv")

    assert manifest["profile"] == "paper"
    assert manifest["dolt_available"] is True
    assert manifest["warmups"] == 2 and manifest["measured_trials"] == 7
    assert audit["passed"] is True and audit["issues"] == []
    assert len(raw) == 333 and len(summary) == 37
    assert all(row["status"] == "ok" and row["correctness"] == "True" for row in raw)
    assert len(sensitivity) == 5
    assert all(row["all_correct"] == "True" for row in sensitivity)

    evaluation = {
        (row["scenario"], row["model"]): row
        for row in summary
        if row["phase"] == "evaluation"
    }
    calibration = [row for row in summary if row["phase"] == "calibration"]
    dolt_versions = sorted(
        {row["dolt_version"] for row in raw if row["model"] == "Dolt"}
    )
    assert dolt_versions == ["dolt version 2.3.1"]
    return manifest, audit, evaluation, calibration, sensitivity_manifest, sensitivity


def pil_font(size: int, *, bold: bool = False, italic: bool = False):
    if bold and italic:
        filename = "timesbi.ttf"
    elif bold:
        filename = "timesbd.ttf"
    elif italic:
        filename = "timesi.ttf"
    else:
        filename = "times.ttf"
    try:
        return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / filename), size)
    except OSError:
        return ImageFont.load_default()


def hex_color(value: str) -> str:
    return value if value.startswith("#") else f"#{value}"


def text_center(draw, box, text, font, fill=INK, spacing=5):
    left, top, right, bottom = box
    bounds = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    draw.multiline_text(
        ((left + right - width) / 2, (top + bottom - height) / 2),
        text,
        font=font,
        fill=hex_color(fill),
        spacing=spacing,
        align="center",
    )


def arrow(draw, start, end, color=INK, width=4):
    color = hex_color(color)
    draw.line((start, end), fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    for delta in (2.55, -2.55):
        point = (
            end[0] + 14 * math.cos(angle + delta),
            end[1] + 14 * math.sin(angle + delta),
        )
        draw.line((end, point), fill=color, width=width)


def make_architecture(path: Path):
    image = Image.new("RGB", (1500, 880), "white")
    draw = ImageDraw.Draw(image)
    title = pil_font(36, bold=True)
    label = pil_font(36, bold=True)
    body = pil_font(32)
    draw.text((55, 35), "Chronos-H architecture", font=title, fill=hex_color(INK))
    layers = [
        ((75, 150, 360, 315), "Frontend\nmetrics workspace", "F3EEF7", "B279A2"),
        ((440, 150, 725, 315), "REST API\nrepository operations", BLUE_LIGHT, BLUE),
        ((805, 150, 1090, 315), "Chronos-H core\ncommit, checkout, diff", "EAF4EA", GREEN),
        ((1170, 150, 1455, 315), "SQLite object store\nobjects, commits, HEAD", "FFF4E6", AMBER),
    ]
    for box, text, fill, outline in layers:
        draw.rounded_rectangle(box, radius=16, fill=hex_color(fill), outline=hex_color(outline), width=4)
        text_center(draw, box, text, label if "Chronos-H" in text else body)
    for left, right in zip(layers, layers[1:]):
        arrow(draw, (left[0][2] + 12, 232), (right[0][0] - 12, 232), BLUE)

    boundary = (280, 430, 1220, 795)
    draw.rounded_rectangle(boundary, radius=20, fill="#FAFAFA", outline=hex_color(BLUE), width=5)
    draw.text((320, 465), "Research contribution boundary", font=label, fill=hex_color(BLUE))
    components = [
        ((330, 570, 570, 720), "Fixed-depth\nMerkle hash trie"),
        ((630, 570, 870, 720), "Incremental\ncopy-on-write hashing"),
        ((930, 570, 1170, 720), "Adaptive log /\nMerkle differencing"),
    ]
    for box, text in components:
        draw.rounded_rectangle(box, radius=13, fill="white", outline=hex_color(BLUE), width=3)
        text_center(draw, box, text, body)
    image.save(path, dpi=(300, 300))


def make_workflow(path: Path):
    image = Image.new("RGB", (1500, 780), "white")
    draw = ImageDraw.Draw(image)
    title = pil_font(36, bold=True)
    body = pil_font(32)
    strong = pil_font(34, bold=True)
    draw.text((55, 35), "Incremental commit and hybrid diff", font=title, fill=hex_color(INK))
    steps = [
        "Canonical\nmutation batch",
        "Hash keys and\ngroup routes",
        "Rewrite affected\nleaf buckets once",
        "Rewrite ancestor\npath union",
        "Persist objects,\ncommit, and HEAD",
    ]
    xs = [55, 345, 635, 925, 1215]
    for index, (x, text) in enumerate(zip(xs, steps)):
        box = (x, 135, x + 225, 285)
        fill = "FFF4E6" if index == 4 else BLUE_LIGHT
        draw.rounded_rectangle(box, radius=14, fill=hex_color(fill), outline=hex_color(BLUE), width=3)
        text_center(draw, box, text, body)
        if index < 4:
            arrow(draw, (x + 235, 210), (xs[index + 1] - 10, 210), BLUE)

    arrow(draw, (750, 310), (750, 405), BLUE)
    draw.text((704, 350), "DIFF", font=strong, fill=hex_color(BLUE))
    left = (260, 470, 675, 690)
    right = (825, 470, 1240, 690)
    draw.rounded_rectangle(left, radius=16, fill="#EAF4EA", outline=hex_color(GREEN), width=4)
    draw.rounded_rectangle(right, radius=16, fill="#FFF4E6", outline=hex_color(AMBER), width=4)
    text_center(draw, left, "Operations <= 4,096\naggregate addressed changesets\n(log path)", body)
    text_center(draw, right, "Operations > 4,096\ncompare unequal subtree hashes\n(Merkle path)", body)
    arrow(draw, (725, 435), (600, 470), GREEN)
    arrow(draw, (775, 435), (900, 470), AMBER)
    image.save(path, dpi=(300, 300))


def log_position(value, minimum, maximum, top, bottom):
    value = max(value, minimum)
    fraction = (math.log10(value) - math.log10(minimum)) / (
        math.log10(maximum) - math.log10(minimum)
    )
    return bottom - fraction * (bottom - top)


def draw_log_axis(draw, left, right, top, bottom, ticks, minimum, maximum, label):
    draw.line((left, top, left, bottom), fill=hex_color(INK), width=3)
    draw.line((left, bottom, right, bottom), fill=hex_color(INK), width=3)
    tick_font = pil_font(40)
    for tick in ticks:
        y = log_position(tick, minimum, maximum, top, bottom)
        draw.line((left, y, right, y), fill="#DDDDDD", width=1)
        text = f"{tick:g}"
        bounds = draw.textbbox((0, 0), text, font=tick_font)
        draw.text((left - 14 - (bounds[2] - bounds[0]), y - 10), text, font=tick_font, fill=hex_color(MUTED))
    draw.text((left, top - 58), label, font=pil_font(42, bold=True), fill=hex_color(INK))


def make_diff_chart(path: Path, evaluation):
    scenarios = ["small-sparse", "medium-sparse", "medium-dense", "large-sparse", "large-hot"]
    models = ["Snapshot", "Log-only", "Chronos-M (forced Merkle)", "Chronos-H", "Dolt"]
    image = Image.new("RGB", (1900, 1080), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 25), "End-to-end version diff latency", font=pil_font(48, bold=True), fill=hex_color(INK))
    draw.text((65, 86), "Median with 25th-75th percentile whiskers; seven measured trials", font=pil_font(36), fill=hex_color(MUTED))
    left, right, top, bottom = 165, 1840, 205, 850
    minimum, maximum = 0.1, 1000
    draw_log_axis(draw, left, right, top, bottom, [0.1, 1, 10, 100, 1000], minimum, maximum, "milliseconds (log scale)")
    group_width = (right - left) / len(scenarios)
    bar_width = 42
    gap = 10
    for group_index, scenario in enumerate(scenarios):
        center = left + group_width * (group_index + 0.5)
        total = len(models) * bar_width + (len(models) - 1) * gap
        start = center - total / 2
        for model_index, model in enumerate(models):
            row = evaluation[(scenario, model)]
            median = float(row["diff_ms_median"])
            low = float(row["diff_ms_p25"])
            high = float(row["diff_ms_p75"])
            x0 = start + model_index * (bar_width + gap)
            x1 = x0 + bar_width
            y = log_position(median, minimum, maximum, top, bottom)
            draw.rectangle((x0, y, x1, bottom), fill=hex_color(COLORS[model]))
            mid_x = (x0 + x1) / 2
            y_low = log_position(low, minimum, maximum, top, bottom)
            y_high = log_position(high, minimum, maximum, top, bottom)
            draw.line((mid_x, y_low, mid_x, y_high), fill=hex_color(INK), width=3)
            draw.line((mid_x - 9, y_low, mid_x + 9, y_low), fill=hex_color(INK), width=3)
            draw.line((mid_x - 9, y_high, mid_x + 9, y_high), fill=hex_color(INK), width=3)
        label = scenario.replace("-", "\n")
        text_center(draw, (center - 140, 870, center + 140, 965), label, pil_font(38))

    legend_y = 1000
    legend_x = 175
    legend_labels = ["Snapshot", "Log-only", "Chronos-M", "Chronos-H", "Dolt"]
    for index, (model, label) in enumerate(zip(models, legend_labels)):
        x = legend_x + index * 320
        draw.rectangle((x, legend_y, x + 35, legend_y + 25), fill=hex_color(COLORS[model]))
        draw.text((x + 48, legend_y - 8), label, font=pil_font(34), fill=hex_color(INK))
    image.save(path, dpi=(300, 300))


def make_operational_chart(path: Path, evaluation):
    scenarios = ["small-sparse", "medium-sparse", "medium-dense", "large-sparse", "large-hot"]
    metrics = [
        ("incremental_commit", "Incremental commit", 10, 2000, [10, 100, 1000]),
        ("checkout", "Historical checkout", 1, 1000, [1, 10, 100, 1000]),
    ]
    image = Image.new("RGB", (1900, 980), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 25), "Chronos-H and Dolt operational latency", font=pil_font(48, bold=True), fill=hex_color(INK))
    draw.text((65, 86), "Median with 25th-75th percentile whiskers; CLI-facing system protocol", font=pil_font(36), fill=hex_color(MUTED))
    panel_width = 790
    panel_gap = 110
    panel_lefts = [155, 155 + panel_width + panel_gap]
    top, bottom = 205, 750
    for panel_index, (metric, title, minimum, maximum, ticks) in enumerate(metrics):
        left = panel_lefts[panel_index]
        right = left + panel_width
        draw_log_axis(draw, left, right, top, bottom, ticks, minimum, maximum, "")
        draw.text((left, 172), "ms (log)", font=pil_font(28, bold=True), fill=hex_color(INK))
        draw.text((left + 220, 138), title, font=pil_font(38, bold=True), fill=hex_color(INK))
        group_width = panel_width / len(scenarios)
        for group_index, scenario in enumerate(scenarios):
            center = left + group_width * (group_index + 0.5)
            for offset, model in [(-27, "Chronos-H"), (27, "Dolt")]:
                row = evaluation[(scenario, model)]
                median = float(row[f"{metric}_ms_median"])
                low = float(row[f"{metric}_ms_p25"])
                high = float(row[f"{metric}_ms_p75"])
                x0, x1 = center + offset - 20, center + offset + 20
                y = log_position(median, minimum, maximum, top, bottom)
                draw.rectangle((x0, y, x1, bottom), fill=hex_color(COLORS[model]))
                y_low = log_position(low, minimum, maximum, top, bottom)
                y_high = log_position(high, minimum, maximum, top, bottom)
                draw.line((center + offset, y_low, center + offset, y_high), fill=hex_color(INK), width=3)
                draw.line((center + offset - 7, y_low, center + offset + 7, y_low), fill=hex_color(INK), width=3)
                draw.line((center + offset - 7, y_high, center + offset + 7, y_high), fill=hex_color(INK), width=3)
            label = {"small-sparse": "S-S", "medium-sparse": "M-S", "medium-dense": "M-D", "large-sparse": "L-S", "large-hot": "L-H"}[scenario]
            bounds = draw.textbbox((0, 0), label, font=pil_font(34))
            draw.text((center - (bounds[2] - bounds[0]) / 2, bottom + 25), label, font=pil_font(34), fill=hex_color(INK))
    draw.rectangle((700, 875, 740, 900), fill=hex_color(COLORS["Chronos-H"]))
    draw.text((752, 865), "Chronos-H", font=pil_font(34), fill=hex_color(INK))
    draw.rectangle((970, 875, 1010, 900), fill=hex_color(COLORS["Dolt"]))
    draw.text((1022, 865), "Dolt", font=pil_font(34), fill=hex_color(INK))
    draw.text((1200, 868), "S/M/L: size; S/D/H: locality", font=pil_font(30), fill=hex_color(MUTED))
    image.save(path, dpi=(300, 300))


def make_calibration_chart(path: Path, calibration):
    by_model = {}
    for row in calibration:
        by_model.setdefault(row["model"], []).append(row)
    for rows in by_model.values():
        rows.sort(key=lambda row: float(row["work_examined_median"]))
    log_rows = by_model["Chronos-log calibration"]
    merkle_rows = by_model["Chronos-M (forced Merkle)"]
    operations = [float(row["work_examined_median"]) for row in log_rows]
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 30), "Independent Chronos-H threshold calibration", font=pil_font(37, bold=True), fill=hex_color(INK))
    draw.text((60, 82), "Calibration workloads are disjoint from final evaluation scenarios", font=pil_font(22), fill=hex_color(MUTED))
    left, right, top, bottom = 170, 1500, 180, 720
    minimum_y, maximum_y = 0.03, 100
    draw_log_axis(draw, left, right, top, bottom, [0.03, 0.1, 1, 10, 100], minimum_y, maximum_y, "diff milliseconds (log scale)")
    minimum_x, maximum_x = math.log10(4), math.log10(4096)

    def xpos(value):
        return left + (math.log10(value) - minimum_x) / (maximum_x - minimum_x) * (right - left)

    for operation in operations:
        x = xpos(operation)
        draw.line((x, bottom, x, bottom + 8), fill=hex_color(INK), width=2)
        label = f"{int(operation):,}"
        bounds = draw.textbbox((0, 0), label, font=pil_font(34))
        draw.text((x - (bounds[2] - bounds[0]) / 2, bottom + 18), label, font=pil_font(34), fill=hex_color(MUTED))
    draw.text((570, 810), "accumulated operations (log scale)", font=pil_font(38, bold=True), fill=hex_color(INK))

    for rows, color, label in [
        (log_rows, GREEN, "Log aggregation"),
        (merkle_rows, AMBER, "Merkle comparison"),
    ]:
        points = []
        for row, operation in zip(rows, operations):
            points.append((xpos(operation), log_position(float(row["diff_ms_median"]), minimum_y, maximum_y, top, bottom)))
        draw.line(points, fill=hex_color(color), width=6)
        for x, y in points:
            draw.ellipse((x - 8, y - 8, x + 8, y + 8), fill=hex_color(color), outline="white", width=2)
        draw.text((1050, 118 if label.startswith("Log") else 160), label, font=pil_font(34, bold=True), fill=hex_color(color))
    threshold_x = xpos(4096)
    draw.line((threshold_x, top, threshold_x, bottom), fill=hex_color(BLUE), width=4)
    draw.text((1040, 610), "Frozen boundary: 4,096", font=pil_font(34, bold=True), fill=hex_color(BLUE))
    image.save(path, dpi=(300, 300))


def make_sensitivity_chart(path: Path, sensitivity):
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 30), "Fixed-trie geometry sensitivity", font=pil_font(37, bold=True), fill=hex_color(INK))
    draw.text((60, 82), "10,000 rows, 100 changes/commit, forced-Merkle diff; seven trials", font=pil_font(22), fill=hex_color(MUTED))
    left, right, top, bottom = 175, 1500, 175, 730
    min_x, max_x = 4.0, 8.7
    min_y, max_y = 15.0, 65.0
    draw.line((left, top, left, bottom), fill=hex_color(INK), width=3)
    draw.line((left, bottom, right, bottom), fill=hex_color(INK), width=3)
    for value in [4, 5, 6, 7, 8]:
        x = left + (value - min_x) / (max_x - min_x) * (right - left)
        draw.line((x, top, x, bottom), fill="#E2E2E2", width=1)
        draw.text((x - 14, bottom + 20), str(value), font=pil_font(34), fill=hex_color(MUTED))
    for value in [20, 30, 40, 50, 60]:
        y = bottom - (value - min_y) / (max_y - min_y) * (bottom - top)
        draw.line((left, y, right, y), fill="#E2E2E2", width=1)
        draw.text((85, y - 17), str(value), font=pil_font(34), fill=hex_color(MUTED))
    draw.text((590, 815), "repository storage (MiB)", font=pil_font(38, bold=True), fill=hex_color(INK))
    draw.text((48, 125), "diff ms", font=pil_font(38, bold=True), fill=hex_color(INK))
    offsets = {"b4-d6": (-110, -52), "b8-d3": (20, -45), "b8-d4": (20, 15), "b8-d5": (20, -18), "b16-d3": (-150, -42)}
    for row in sensitivity:
        storage = float(row["storage_bytes_median"]) / (1024 * 1024)
        diff = float(row["diff_ms_median"])
        x = left + (storage - min_x) / (max_x - min_x) * (right - left)
        y = bottom - (diff - min_y) / (max_y - min_y) * (bottom - top)
        config = row["config"]
        color = SENSITIVITY_COLORS[config]
        radius = 13 if config == "b8-d4" else 10
        draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=hex_color(color))
        dx, dy = offsets[config]
        draw.text((x + dx, y + dy), config, font=pil_font(34, bold=config == "b8-d4"), fill=hex_color(color))
    draw.text((180, 765), "Lower-left is better; the current b8-d4 default is highlighted.", font=pil_font(32, italic=True), fill=hex_color(MUTED))
    image.save(path, dpi=(300, 300))


def set_font(run, size=None, *, bold=None, italic=None, name="Times New Roman", color=INK):
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), name)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_columns(section, count: int):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    if count == 2:
        cols.set(qn("w:space"), "346")  # 0.24 in / 6 mm, per the BERT formatting guide
        cols.set(qn("w:equalWidth"), "1")
    else:
        cols.attrib.pop(qn("w:space"), None)
        cols.attrib.pop(qn("w:equalWidth"), None)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)
    set_columns(section, 1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = Pt(13.6)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.15)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(14.5)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(28)
    title.paragraph_format.line_spacing = Pt(16)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in [
        ("Heading 1", 12),
        ("Heading 2", 11),
        ("Heading 3", 10),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = Pt(13.6)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10)
    caption.font.bold = False
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(INK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.line_spacing = Pt(12)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.2)
        style.paragraph_format.first_line_indent = Inches(-0.15)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = Pt(13.6)

    if "IEEE Reference" not in [style.name for style in styles]:
        reference = styles.add_style("IEEE Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        reference = styles["IEEE Reference"]
    reference.font.name = "Times New Roman"
    reference.font.size = Pt(10)
    reference.paragraph_format.left_indent = Inches(0.15)
    reference.paragraph_format.first_line_indent = Inches(-0.15)
    reference.paragraph_format.space_after = Pt(2)
    reference.paragraph_format.line_spacing = Pt(10.5)

    for sec in doc.sections:
        sec.header.paragraphs[0].text = ""
        sec.footer.paragraphs[0].text = ""

    settings = doc.settings._element
    auto_hyphenation = settings.find(qn("w:autoHyphenation"))
    if auto_hyphenation is None:
        auto_hyphenation = OxmlElement("w:autoHyphenation")
        settings.append(auto_hyphenation)
    auto_hyphenation.set(qn("w:val"), "true")


def enforce_submission_typography(doc: Document):
    """Keep every visible run in the reference paper's black serif family."""
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    for paragraph in paragraphs:
        for run in paragraph.runs:
            set_font(run, color=INK)


def normalize_heading_following_indents(doc: Document):
    """Match ACL convention: the first prose paragraph after a heading is flush left."""
    after_heading = False
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            after_heading = True
            continue
        if not paragraph.text.strip():
            continue
        if after_heading and paragraph.style.name == "Normal":
            paragraph.paragraph_format.first_line_indent = Inches(0)
        after_heading = False


def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def link_citations(doc: Document, citation_targets):
    """Turn numeric citations into blue internal links to numbered references."""
    ordered = sorted(citation_targets, key=len, reverse=True)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "IEEE Reference":
            continue
        for run in list(paragraph.runs):
            if not any(citation in run.text for citation in ordered):
                continue
            source = run.text
            pieces = []
            while source:
                matches = [
                    (source.find(citation), citation)
                    for citation in ordered
                    if source.find(citation) >= 0
                ]
                if not matches:
                    pieces.append((source, None))
                    break
                index, citation = min(matches, key=lambda item: item[0])
                if index:
                    pieces.append((source[:index], None))
                pieces.append((citation, citation_targets[citation]))
                source = source[index + len(citation) :]

            parent = run._r.getparent()
            insert_at = parent.index(run._r)
            parent.remove(run._r)
            for offset, (text, anchor) in enumerate(pieces):
                if not text:
                    continue
                new_run = deepcopy(run._r)
                for child in list(new_run):
                    if child.tag != qn("w:rPr"):
                        new_run.remove(child)
                text_node = OxmlElement("w:t")
                if text[0].isspace() or text[-1].isspace():
                    text_node.set(qn("xml:space"), "preserve")
                text_node.text = text
                new_run.append(text_node)
                node = new_run
                if anchor:
                    run_properties = new_run.get_or_add_rPr()
                    color = ensure_ooxml_child(run_properties, "w:color")
                    color.set(qn("w:val"), HYPERLINK_BLUE)
                    underline = ensure_ooxml_child(run_properties, "w:u")
                    underline.set(qn("w:val"), "none")
                    hyperlink = OxmlElement("w:hyperlink")
                    hyperlink.set(qn("w:anchor"), anchor)
                    hyperlink.set(qn("w:history"), "1")
                    hyperlink.append(new_run)
                    node = hyperlink
                parent.insert(insert_at + offset, node)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, *, edges=(), color=GRID, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single" if edge in edges else "nil")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row(row, repeat=False):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def add_table(doc, headers, rows, weights, *, font_size=9, width_dxa=4270):
    font_size = 9
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    header = table.rows[0]
    keep_row(header, repeat=True)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = str(value)
        set_cell_shading(cell, WHITE)
        set_cell_border(cell, edges=("top", "bottom"), size="6")
        set_cell_margins(cell, top=20, start=25, bottom=20, end=25)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_font(run, font_size, bold=True)
    for row_values in rows:
        row = table.add_row()
        keep_row(row)
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.text = str(value)
            set_cell_border(cell)
            set_cell_margins(cell, top=16, start=25, bottom=16, end=25)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(10)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_font(run, font_size)
    for cell in table.rows[-1].cells:
        set_cell_border(cell, edges=("bottom",), size="6")
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    widths = column_widths_from_weights(weights, width_dxa)
    apply_table_geometry(table, widths, table_width_dxa=width_dxa, indent_dxa=25)
    return table


def add_body(doc, text, *, indent=True):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.15 if indent else 0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(13.6)
    paragraph.add_run(text)
    return paragraph


def add_section(doc, title):
    return doc.add_heading(title, level=1)


def add_subsection(doc, title):
    return doc.add_heading(title, level=2)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_picture(doc, path, alt_text, *, width=Inches(2.96)):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text.split(".")[0])
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.paragraph_format.keep_with_next = False
    paragraph.add_run(text)
    return paragraph


def add_source_note(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_font(run, 9, italic=True, color=MUTED)
    return paragraph


def add_two_column_section(doc):
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)
    set_columns(section, 2)
    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""
    return section


def build():
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    manifest, audit, evaluation, calibration, sensitivity_manifest, sensitivity = load_evidence()

    architecture_path = ASSET_DIR / "architecture.png"
    workflow_path = ASSET_DIR / "workflow.png"
    diff_path = ASSET_DIR / "diff_results.png"
    operational_path = ASSET_DIR / "operational_results.png"
    calibration_path = ASSET_DIR / "calibration.png"
    sensitivity_path = ASSET_DIR / "sensitivity.png"
    make_architecture(architecture_path)
    make_workflow(workflow_path)
    make_diff_chart(diff_path, evaluation)
    make_operational_chart(operational_path, evaluation)
    make_calibration_chart(calibration_path, calibration)
    make_sensitivity_chart(sensitivity_path, sensitivity)

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Chronos: A Content-Addressed Versioned Data Store"
    doc.core_properties.subject = "Final research paper"
    doc.core_properties.author = "Adhyan Jain; Atharva Sheersh Pandey"
    doc.core_properties.keywords = (
        "content-addressed storage, structured data versioning, Merkle trie, "
        "incremental hashing, adaptive diff, Dolt"
    )

    title = doc.add_paragraph(style="Title")
    title.add_run(
        "Chronos: A Content-Addressed Versioned Data Store: A Git-Inspired Approach to "
        "Structured Data Versioning Using Merkle Search Trees"
    )
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.first_line_indent = Inches(0)
    author.paragraph_format.space_after = Pt(0)
    run = author.add_run("[Adhyan Jain]")
    set_font(run, 12)
    marker = author.add_run("1")
    set_font(marker, 8)
    marker.font.superscript = True
    run = author.add_run(", [Atharva Sheersh Pandey]")
    set_font(run, 12)
    marker = author.add_run("1")
    set_font(marker, 8)
    marker.font.superscript = True
    affiliation = doc.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation.paragraph_format.first_line_indent = Inches(0)
    affiliation.paragraph_format.space_after = Pt(0)
    marker = affiliation.add_run("1")
    set_font(marker, 8)
    marker.font.superscript = True
    run = affiliation.add_run("Department of Computer Science and Engineering, VIT Vellore")
    set_font(run, 11)
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.first_line_indent = Inches(0)
    contact.paragraph_format.space_after = Pt(0)
    run = contact.add_run("adhyan.jain2024@vitstudent.ac.in")
    set_font(run, 10.5, color=HYPERLINK_BLUE)
    run.font.underline = True
    run = contact.add_run(", atharva.sheersh2024@vitstudent.ac.in")
    set_font(run, 10.5, color=HYPERLINK_BLUE)
    run.font.underline = True
    guide = doc.add_paragraph()
    guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    guide.paragraph_format.first_line_indent = Inches(0)
    guide.paragraph_format.space_after = Pt(34)
    run = guide.add_run("Guide: Dr. Poornima N, VIT Vellore")
    set_font(run, 11)

    add_two_column_section(doc)

    abstract_heading = doc.add_paragraph()
    abstract_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abstract_heading.paragraph_format.first_line_indent = Inches(0)
    abstract_heading.paragraph_format.space_after = Pt(7)
    lead = abstract_heading.add_run("Abstract")
    set_font(lead, 12, bold=True)
    abstract = doc.add_paragraph()
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.first_line_indent = Inches(0)
    abstract.paragraph_format.left_indent = Pt(17)
    abstract.paragraph_format.right_indent = Pt(17)
    abstract.paragraph_format.space_after = Pt(4)
    abstract.paragraph_format.line_spacing = Pt(12)
    body = abstract.add_run(
        "Versioning structured datasets requires durable history, efficient sparse updates, "
        "historical reconstruction, and comparison without copying or scanning every record. "
        "Chronos is a Python prototype built around an immutable fixed-depth Merkle hash trie. "
        "A commit rewrites only affected leaf buckets and the union of their ancestor paths; "
        "unchanged subtrees are reused by content hash. Chronos-H adds addressed changesets and "
        "selects between log aggregation and hash-pruned tree differencing using a separately "
        "calibrated threshold. We evaluate Snapshot, Log-only, forced-Merkle Chronos-M, "
        "Chronos-H, and Dolt 2.3.1 on identical deterministic workloads up to 100,000 rows. "
        "The audited paper run contains 333 successful and correct trials. Under the CLI-facing "
        "protocol, Chronos-H diff latency was 5.8x to 554.2x lower than Dolt and its incremental "
        "commit latency was 2.3x to 37.4x lower, while Chronos-H consumed 1.83x to 6.43x more "
        "repository storage. Against forced Merkle, the hybrid log path reduced median sparse or "
        "hot diff latency by 6.0x to 10.0x. A five-configuration sensitivity study shows that "
        "the current b=8, d=4 trie minimizes storage in the tested set, whereas b=8, d=3 favors "
        "latency. The results support a workload-dependent hybrid design, not universal "
        "superiority over a production SQL system."
    )
    set_font(body, 10)

    add_section(doc, "1  Introduction")
    add_body(
        doc,
        "Version control for structured data is more than archiving files. A useful system must "
        "retain exact versions, apply sparse mutations without rebuilding the complete state, "
        "reconstruct historical versions, and identify changes efficiently. Git provides the "
        "starting idea--immutable objects named by content and commits linked to their parents "
        "[7]--but organizes those objects around files and directories. A record store presents "
        "a different workload: keys remain stable, update density varies, and two versions may "
        "be separated by a handful of edits or by a long sequence of commits.",
        indent=False,
    )
    add_body(
        doc,
        "Chronos evaluates whether a deliberately simple fixed-depth Merkle hash trie can unite "
        "incremental structural sharing with an operation index. Its primary variant, Chronos-H, "
        "uses the log path for short or sparse intervals and switches to the Merkle path when the "
        "accumulated operations cross a calibration boundary fixed before evaluation. SQLite "
        "provides atomic durable storage, but does not determine the versioning algorithm."
    )
    add_body(
        doc,
        "Dolt is the production comparator because it combines SQL semantics, a Git-style commit "
        "graph, and content-addressed Prolly Trees [8], [9]. Chronos is not presented as a feature "
        "replacement for Dolt. The comparison instead identifies the performance and storage "
        "trade-offs of a smaller instrumentable mechanism under a reproducible CLI-facing protocol."
    )

    add_subsection(doc, "1.1  Research Questions")
    add_bullets(
        doc,
        [
            "RQ1: What latency and storage trade-offs does Chronos-H exhibit against Snapshot, Log-only, and Dolt across dataset size, mutation density, and locality?",
            "RQ2: Does adaptive log/Merkle selection improve diff latency relative to a forced-Merkle ablation without changing correctness?",
            "RQ3: How sensitive are fixed-trie latency, storage, and examined work to branching factor and depth?",
        ],
    )
    add_subsection(doc, "1.2  Contributions")
    add_bullets(
        doc,
        [
            "A canonical SHA-256-addressed fixed-depth trie with incremental multi-key path rewriting.",
            "A hybrid diff policy combining addressed changesets with hash-pruned tree comparison.",
            "An atomic SQLite repository that persists objects, commits, aliases, metrics, and HEAD and verifies hashes after reopen.",
            "A five-variant benchmark with shared deterministic workloads, correctness oracles, Dolt integration, raw evidence, and an independently audited summary.",
            "A controlled branching-factor/depth sensitivity study that prevents treating b=8, d=4 as universally optimal.",
        ],
    )

    add_section(doc, "2  Literature Review and Related Work")
    add_subsection(doc, "2.1  Content Addressing and Persistent Structures")
    add_body(
        doc,
        "Merkle's authenticated tree construction established that a compact root digest can "
        "commit to a larger collection [13], while persistent-data-structure theory formalized "
        "path copying and bounded-overhead access to historical versions [10]. Git applies "
        "content-addressed blobs, trees, and commits to software history [7], and IPFS generalizes "
        "content-addressed links into a versionable Merkle DAG [4]. Chronos adopts immutable object "
        "identity and parent links, but routes structured keys by hash into a persistent trie rather "
        "than mirroring a filesystem hierarchy.",
        indent=False,
    )
    add_body(
        doc,
        "Merkle Search Trees combine search-tree ordering with Merkle authentication for efficient "
        "state reconciliation [3]. Chronos shares hash-pruned comparison but intentionally uses "
        "fixed-depth routing and a centralized commit history rather than a replicated CRDT. "
        "Content-defined chunking, established in LBFS for redundancy detection [14], also explains "
        "why history-independent boundaries are a credible storage improvement beyond the current "
        "fixed buckets."
    )
    add_subsection(doc, "2.2  Structured Dataset Versioning")
    add_body(
        doc,
        "Dataset-versioning research exposes a broader storage-retrieval design space. DataHub "
        "motivates collaborative dataset history at scale [5]. Decibel integrates branching into "
        "a relational storage engine and compares version-first, tuple-first, and hybrid layouts "
        "[12]. OrpheusDB bolts versioning onto a conventional DBMS and optimizes partitioning for "
        "version retrieval [11]. The dataset-versioning model in [6] formalizes the conflict between "
        "storing more materialized state and paying more reconstruction cost. Delta Lake represents "
        "the complementary log-and-checkpoint design for ACID tables and time travel [1]. These "
        "systems motivate "
        "reporting commit, checkout, diff, and storage together rather than selecting a single metric."
    )
    add_subsection(doc, "2.3  System Design Gap")
    add_body(
        doc,
        "Noms represents structured data as a Merkle DAG of immutable chunks and supports "
        "efficient diff and synchronization [2]. ForkBase combines content addressing, fork "
        "semantics, and duplicate-content detection for forkable applications [15]. Dolt stores "
        "table indexes as content-addressed Prolly Trees: ordered, B-tree-like structures with "
        "history-independent chunk boundaries, structural sharing, and native diff [8], [9]. "
        "Chronos differs by using fixed-depth hash routing and by exposing an explicit alternative "
        "operation-log path. This simplifies instrumentation, but gives up Dolt's ordered range "
        "behavior, SQL surface, mature branching, merging, and production engineering."
    )
    add_table(
        doc,
        ["Prior work", "Core representation", "Relationship to Chronos"],
        [
            ("Persistent DS [10]", "Path copying", "Formal structural-sharing basis"),
            ("Git [7]", "Object DAG", "Identity and commit ancestry; file-oriented"),
            ("IPFS [4]", "Merkle DAG", "General content-addressed versioning"),
            ("MST [3]", "Merkle search tree", "Authenticated ordered reconciliation"),
            ("DataHub [5]", "Version graph", "Collaborative dataset motivation"),
            ("Decibel [12]", "Relational layouts", "Native database branching"),
            ("OrpheusDB [11]", "Partitioned versions", "Storage/retrieval optimization"),
            ("Delta Lake [1]", "Log/checkpoints", "ACID time travel at data-lake scale"),
            ("Noms [2]", "Merkle DAG", "Structured content addressing"),
            ("ForkBase [15]", "Forkable CAS", "Fork and deduplication semantics"),
            ("DoltHub [8,9]", "Prolly Trees", "Production SOTA comparator"),
        ],
        [0.85, 1.2, 1.45],
        font_size=6.8,
    )
    add_caption(doc, "Table 1: Literature-review matrix and the gap addressed by Chronos.")

    add_subsection(doc, "2.4  Design Implications and Research Positioning")
    add_body(
        doc,
        "Three separate costs recur in the literature: retained history, creation of a new version, "
        "and later reconstruction or comparison. A snapshot keeps reconstruction straightforward "
        "by copying unchanged records. A log avoids that copy at commit time, but the deferred work "
        "returns during replay, compaction, or checkpointing. Path-copying persistent structures "
        "take a third route: only the paths changed by an update are new, while the remaining "
        "topology is shared [10]. The storage-versus-recreation model in [6] makes this trade-off "
        "explicit. Decibel and OrpheusDB likewise show that one physical layout does not dominate "
        "scans, version retrieval, and storage for every branching workload [11], [12]. For that "
        "reason, our evaluation reports import, commit, "
        "checkout, diff, and storage rather than using one measurement as a proxy for the system."
    )
    add_body(
        doc,
        "Change detection is another point of separation. Comparing two complete states does not "
        "need a detailed history, although its work grows with the states being scanned. Replaying "
        "changesets instead makes the interval length decisive and requires an unbroken ancestry "
        "path with complete changesets. With Merkle differencing, matching subtree hashes stop the search; only mismatching "
        "branches are opened. The observed cost then depends on trie shape, leaf occupancy, key "
        "distribution, and where updates land. The MST retains key order [3]. Noms and Dolt use "
        "content-defined boundaries to form history-independent Prolly Trees [2], [9], while "
        "ForkBase reuses addressed content across objects, branches, and versions [15]. Chronos chooses fixed "
        "hash-derived routes. This choice removes ordered range access and adaptive boundaries, but "
        "gives the experiment a stable geometry whose depth and branching factor can be measured."
    )
    add_body(
        doc,
        "The gap addressed here is not a new content-addressed store, but a controlled study in "
        "which two correct diff algorithms share one durable history and are chosen from observable "
        "properties of the requested interval. Every Chronos-H commit records both a Merkle root and an addressed changeset, so "
        "changing the diff path does not change the meaning of a version. We expect log aggregation "
        "to suit short, sparse intervals and hash pruning to become preferable after many operations. "
        "We also expect b and d to change latency, examined work, and storage. Snapshot and Log-only provide "
        "representation baselines; Chronos-M removes the selector; Chronos-H retains it; and Dolt "
        "provides a production reference point. This design tests the hybrid mechanism while keeping "
        "the system-level Dolt comparison separate from any claim about intrinsic tree superiority."
    )

    add_section(doc, "3  Chronos Design")
    add_picture(
        doc,
        architecture_path,
        "Chronos layered architecture showing the frontend, REST API, Chronos-H core, SQLite object store, and the research contribution boundary.",
    )
    add_caption(doc, "Figure 1: Chronos system architecture and the boundary of its research contribution.")
    add_subsection(doc, "3.1  Addressed Object Model")
    add_body(
        doc,
        "Chronos serializes trie nodes and changesets as canonical JSON, then names each object with "
        "its SHA-256 digest. The commit digest includes the root, parent, changeset, message, and "
        "timestamp. Labels such as v1 and v2 are only human-readable aliases; they do not determine "
        "object identity. The durable HEAD reference stores the current commit.",
        indent=False,
    )
    add_subsection(doc, "3.2  Fixed-Depth Merkle Hash Trie")
    add_body(
        doc,
        "For the main experiment we set b=8 and d=4. Each level reads three bits from SHA-256(key), "
        "so twelve routing bits address 4,096 possible leaf buckets. A bucket still stores the full "
        "key, which means that two keys sharing a routing prefix remain distinguishable. Under "
        "uniform hashes the expected bucket occupancy is N/b^d. We test other b,d pairs separately "
        "instead of assuming that this default is optimal."
    )
    add_subsection(doc, "3.3  Incremental Hashing and Commit")
    add_body(
        doc,
        "A mutation batch is canonicalized, routed, and grouped by shared prefixes. Each affected "
        "leaf is rewritten once, and only the union of affected ancestor paths is rebuilt. "
        "Unchanged child hashes are copied into new internal nodes, providing structural sharing. "
        "One SQLite transaction then inserts new objects, the changeset, commit metadata, version "
        "alias, measurements, and updated HEAD. Reopen verifies payload hashes and references."
    )
    add_picture(
        doc,
        workflow_path,
        "Incremental commit workflow followed by Chronos-H selection between addressed changeset aggregation and hash-pruned Merkle comparison.",
    )
    add_caption(doc, "Figure 2: Incremental commit path and adaptive diff decision.")
    add_subsection(doc, "3.4  Hybrid Diff")
    add_body(
        doc,
        "The log path walks the ancestor chain and aggregates key operations. The Merkle path "
        "recursively compares roots, prunes equal subtree hashes, and examines full leaf entries "
        "only below unequal paths. Chronos-H counts accumulated ancestor operations and uses the "
        "log at or below 4,096 operations; it uses Merkle above that boundary. Chronos-M forces "
        "the Merkle path and therefore serves only as an ablation."
    )

    add_section(doc, "4  Experimental Methodology")
    add_subsection(doc, "4.1  Variants and Workloads")
    add_body(
        doc,
        "Five implementations receive identical logical states and mutation batches: a complete "
        "Snapshot per version, an initial state plus operation-log baseline, forced-Merkle "
        "Chronos-M, adaptive Chronos-H, and Dolt 2.3.1 using a keyed SQL table and native diff. "
        "Chronos-M is not a headline product variant; it isolates the selector.",
        indent=False,
    )
    add_table(
        doc,
        ["Scenario", "Rows", "Comm.", "Changes", "Locality"],
        [
            ("Small sparse", "1,000", "10", "10", "spread"),
            ("Medium sparse", "10,000", "10", "10", "spread"),
            ("Medium dense", "10,000", "10", "1,000", "spread"),
            ("Large sparse", "100,000", "10", "100", "spread"),
            ("Large hot", "100,000", "10", "100", "hot"),
        ],
        [1.2, 0.75, 0.72, 0.85, 0.9],
        font_size=6.7,
    )
    add_caption(doc, "Table 2: Final evaluation workload matrix.")
    add_subsection(doc, "4.2  Measurement Contract")
    add_body(
        doc,
        "Initial import measures creation of the first durable state; incremental commit is the "
        "within-trial median across ten durable commits; diff consumes the complete initial-to-final "
        "comparison; checkout fully materializes the target version; storage is the sum of regular "
        "repository-file bytes. Correctness requires exact initial and final states plus a changed-key "
        "oracle. Work examined uses each system's natural unit and is not normalized across keys, "
        "operations, and trie-node pairs."
    )
    add_subsection(doc, "4.3  Repetition, Calibration, and Integrity")
    add_body(
        doc,
        f"Before measurement, each configuration runs twice without contributing samples. The next "
        f"seven trials are recorded, using base seed {manifest['base_seed']}. For every adapter we "
        f"regenerate the workload and compare its SHA-256 digest with the manifest before execution. "
        f"Calibration uses a separate workload set. Reported medians and quartiles exclude warmups. "
        f"The recorded run identifier is {manifest['run_id']}; it was produced on "
        f"{manifest['platform']} with Python {manifest['python_version']}. A separate audit rebuilt "
        f"the trial matrix from the raw CSV and recalculated every summary value."
    )
    add_picture(
        doc,
        calibration_path,
        "Calibration plot comparing log aggregation and forced-Merkle diff latency across accumulated operation counts, with the frozen 4,096-operation boundary.",
    )
    add_caption(doc, "Figure 3: Independent threshold calibration; 4,096 is the largest sampled log-favorable point.")
    add_body(
        doc,
        "The boundary is a frozen policy for this machine, not proof that the true crossover is "
        "exactly 4,096. All 333 paper-profile rows completed with status ok and correctness true. "
        "The audit retained all 83 Tukey outlier candidates; no observation was manually deleted."
    )

    add_section(doc, "5  Results")
    add_subsection(doc, "5.1  Final Diff Performance")
    add_picture(
        doc,
        diff_path,
        "Log-scale grouped bar chart of final median diff latency with interquartile whiskers for Snapshot, Log-only, Chronos-M, Chronos-H, and Dolt across five workloads.",
        width=Inches(2.82),
    )
    add_caption(doc, "Figure 4: Final diff latency from the audited CSV bundle; lower is better.")
    result_rows = []
    for scenario, label in [
        ("small-sparse", "Small sparse"),
        ("medium-sparse", "Medium sparse"),
        ("medium-dense", "Medium dense"),
        ("large-sparse", "Large sparse"),
        ("large-hot", "Large hot"),
    ]:
        h = evaluation[(scenario, "Chronos-H")]
        d = evaluation[(scenario, "Dolt")]
        speedup = float(d["diff_ms_median"]) / float(h["diff_ms_median"])
        result_rows.append(
            (
                label,
                h["strategy_selected"],
                f"{float(h['diff_ms_median']):.3f}",
                f"{float(d['diff_ms_median']):.2f}",
                f"{speedup:.1f}x",
            )
        )
    add_table(
        doc,
        ["Scenario", "Path", "H ms", "Dolt ms", "Ratio"],
        result_rows,
        [1.25, 0.7, 0.72, 0.85, 0.68],
        font_size=6.7,
    )
    add_caption(doc, "Table 3: Chronos-H versus Dolt median diff latency.")
    add_body(
        doc,
        "Chronos-H selected log for the four sparse or hot workloads and Merkle for the 10,000-operation "
        "medium-dense workload. Relative to Dolt, its median diff was 5.8x lower in medium-dense and "
        "58.7x to 554.2x lower in the other workloads. These ratios include repeated Dolt CLI and SQL "
        "invocation costs and therefore describe the evaluated system protocol, not isolated Prolly "
        "Tree algorithm speed."
    )
    add_subsection(doc, "5.2  Hybrid Ablation")
    add_body(
        doc,
        "For log-selected workloads, Chronos-H reduced median diff latency relative to forced-Merkle "
        "Chronos-M by 6.0x (small sparse), 8.7x (medium sparse), 8.8x (large sparse), and 10.0x "
        "(large hot). In medium-dense, both used Merkle and Chronos-H was within 7.3% of the forced "
        "variant. This supports the selector's intended mechanism without claiming that its threshold "
        "is hardware-independent."
    )
    add_subsection(doc, "5.3  Commit, Checkout, and Storage")
    add_picture(
        doc,
        operational_path,
        "Two-panel log-scale chart comparing Chronos-H and Dolt median incremental commit and historical checkout latency with interquartile whiskers.",
    )
    add_caption(doc, "Figure 5: Chronos-H and Dolt operational latency; lower is better.")
    add_body(
        doc,
        "Chronos-H incremental commit medians were 2.3x to 37.4x lower than Dolt, and checkout "
        "medians were 35.8x to 123.4x lower. Initial import was mixed: Chronos-H was 4.9x faster "
        "for small sparse and 1.14x faster for medium sparse, approximately tied for medium dense, "
        "and 2.73x to 2.79x slower at 100,000 rows. The latency advantage is therefore not uniform."
    )
    add_body(
        doc,
        "Chronos-H paid a consistent storage cost. Its median repository footprint was 1.83x, "
        "3.16x, 2.19x, 4.10x, and 6.43x Dolt's for small sparse through large hot, respectively. "
        "The 100,000-row Chronos repositories occupied about 34.4 MiB. This reflects inline values, "
        "JSON object encoding, and the absence of compact chunk packing and garbage collection."
    )
    add_subsection(doc, "5.4  Trie-Parameter Sensitivity")
    add_picture(
        doc,
        sensitivity_path,
        "Scatter plot of forced-Merkle median diff latency against repository storage for five branching-factor and depth configurations, highlighting b8-d4.",
    )
    add_caption(doc, "Figure 6: Fixed-trie sensitivity; the preferred point depends on the objective.")
    add_body(
        doc,
        "All 45 sensitivity trials were correct. The b8-d3 configuration delivered the lowest "
        "median diff latency (20.15 ms versus 44.43 ms for b8-d4) and compared 84.3% fewer node "
        "pairs, but consumed 61.2% more storage and examined three times as many leaf entries. "
        "The current b8-d4 default used the least storage (4.27 MiB). The over-deep b8-d5 design "
        "was slower for every timed operation and used 44.8% more storage. Thus b8-d4 is a "
        "storage-conscious balanced default, not a universal optimum."
    )

    add_section(doc, "6  Discussion")
    add_body(
        doc,
        "The results answer RQ1 with a trade-off rather than a single winner. Chronos-H offered "
        "low diff, commit, and checkout latency under the benchmark interface, but its import "
        "advantage disappeared at 100,000 rows and its repository was always larger than Dolt's. "
        "Snapshot and Log-only also remain relevant baselines: their simpler representations often "
        "reduced import or storage cost. Snapshot comparison, however, grew with the number of "
        "materialized keys.",
        indent=False,
    )
    add_body(
        doc,
        "RQ2 is supported more directly. Because Chronos-H and Chronos-M share the same durable "
        "trie, their sparse-workload diff gap isolates adaptive selection more cleanly than an "
        "absolute comparison with Dolt. The log path avoided thousands of trie comparisons when "
        "only 100 or 1,000 operations separated the endpoints; the dense case switched to Merkle."
    )
    add_body(
        doc,
        "RQ3 rejects a universal geometry. Shallow routing lowers internal traversal but creates "
        "larger leaf buckets and higher storage; deeper routing lowers occupancy but increases "
        "internal-node work. The suitable configuration therefore depends on which requirement "
        "dominates the workload: latency, storage, or worst-case leaf-bucket scanning."
    )
    add_body(
        doc,
        "The Dolt ratios need a system-level interpretation. Dolt performs SQL parsing, schema "
        "management, version-control operations, and production durability work that the Chronos "
        "prototype does not. Dolt's Prolly Trees also serve ordered queries, branches, merges, and "
        "remote workflows. Our timings compare the declared interfaces used in the experiment. They "
        "do not isolate the tree algorithms and therefore cannot show that a fixed trie is inherently "
        "faster than a Prolly Tree."
    )

    add_section(doc, "7  Limitations and Threats to Validity")
    add_bullets(
        doc,
        [
            "Runtime and interface asymmetry: Chronos executes in-process in Python, whereas Dolt is a mature Go executable invoked repeatedly through CLI and SQL interfaces.",
            "Feature asymmetry: Chronos supports one writer and a linear history, but not SQL, branches, merges, remotes, schema evolution, or production concurrency.",
            "Synthetic scope: the deterministic key/value workloads reach 100,000 rows, but cannot represent every real schema, key skew, payload size, or history length.",
            "Single environment: all results were obtained on one Windows 11 machine, where cache state, scheduling, filesystem behavior, and antivirus activity can affect short measurements.",
            "Calibration boundary: 4,096 is the largest sampled operation count that favored the log; this does not establish a universal crossover.",
            "Memory comparability: Python tracemalloc and Dolt process-memory readings are not equivalent, so no cross-system memory graph is reported.",
            "Work-unit mismatch: keys, log operations, trie-node pairs, and Dolt's internal work do not form a single normalized measure.",
            "Outlier policy: all 83 Tukey candidates were kept. Medians and interquartile ranges limit their influence, but seven measured trials do not demonstrate normality.",
        ],
    )

    add_section(doc, "8  Future Work")
    add_body(
        doc,
        "A next version should choose its structure from observed workload statistics. Key count, "
        "prefix skew, mutation density, and leaf occupancy could guide the choice among measured "
        "trie geometries instead of fixing b and d for every repository. The diff selector could "
        "recalibrate periodically and consider history distance, changed-prefix density, and observed "
        "I/O cost alongside its present operation-count rule.",
        indent=False,
    )
    add_body(
        doc,
        "Storage efficiency could be improved with binary node records, separate value chunks, "
        "packed objects, compression, reachability checks, and garbage "
        "collection. An ordered Merkle index or content-defined chunk boundaries could add range "
        "scans and a history-independent layout while leaving the fixed trie as the controlled baseline."
    )
    add_body(
        doc,
        "To move beyond a linear prototype, Chronos still needs named branches, merge commits, "
        "conflict reporting, schema-aware records, concurrent readers and writers, and remote synchronization. "
        "Evaluation should add "
        "public structured datasets, one-million-row and longer-history workloads, Linux replication, "
        "and unified external process telemetry for time, CPU, I/O, and peak resident memory."
    )

    conclusion_column = doc.add_paragraph()
    conclusion_column.paragraph_format.space_after = Pt(0)
    conclusion_column.add_run().add_break(WD_BREAK.PAGE)
    add_section(doc, "9  Conclusion")
    add_body(
        doc,
        "Chronos demonstrates a complete content-addressed versioning path for structured key/value "
        "data: canonical immutable objects, incremental copy-on-write Merkle trie updates, atomic "
        "SQLite persistence, historical checkout, and adaptive differencing. The audited 333-trial "
        "comparison shows that Chronos-H can substantially reduce CLI-facing diff, commit, and "
        "checkout latency relative to Dolt 2.3.1, but it also exposes slower large imports, higher "
        "storage consumption, a narrower feature set, and runtime/interface asymmetry. The ablation "
        "supports hybrid log/Merkle selection, while the sensitivity experiment shows that trie "
        "geometry must be chosen by workload objective. The defensible contribution is therefore "
        "an instrumentable hybrid design and reproducible evidence of its trade-offs, not a claim of "
        "universal replacement for a production version-controlled SQL database.",
        indent=False,
    )

    add_section(doc, "Acknowledgment")
    add_body(
        doc,
        "The authors acknowledge the open-source Dolt project and the researchers whose dataset-versioning systems informed the experimental design.",
        indent=False,
    )

    add_section(doc, "References")
    references = [
        ("ref_delta", "[1] Michael Armbrust et al. 2020. Delta Lake: High-Performance ACID Table Storage over Cloud Object Stores. Proceedings of the VLDB Endowment, 13(12):3411-3424. doi:10.14778/3415478.3415560."),
        ("ref_noms", "[2] Attic Labs. n.d. Noms technical overview. GitHub. https://github.com/attic-labs/noms/blob/master/doc/intro.md."),
        ("ref_mst", "[3] Alex Auvolat and Francois Taiani. 2019. Merkle Search Trees: Efficient State-Based CRDTs in Open Networks. 38th IEEE International Symposium on Reliable Distributed Systems, pages 1-10. doi:10.1109/SRDS47363.2019.00032."),
        ("ref_ipfs", "[4] Juan Benet. 2014. IPFS - Content Addressed, Versioned, P2P File System. arXiv:1407.3561."),
        ("ref_datahub", "[5] Anant Bhardwaj, Souvik Bhattacherjee, Amit Chavan, Amol Deshpande, Aaron J. Elmore, Samuel Madden, and Aditya G. Parameswaran. 2015. DataHub: Collaborative Data Science and Dataset Version Management at Scale. 7th Biennial Conference on Innovative Data Systems Research."),
        ("ref_versioning", "[6] Souvik Bhattacherjee, Amit Chavan, Silu Huang, Amol Deshpande, and Aditya Parameswaran. 2015. Principles of Dataset Versioning: Exploring the Recreation/Storage Tradeoff. Proceedings of the VLDB Endowment, 8(12):1346-1357. doi:10.14778/2824032.2824035."),
        ("ref_git", "[7] Scott Chacon and Ben Straub. 2014. Git Internals - Git Objects. In Pro Git, 2nd edition. https://git-scm.com/book/en/v2/Git-Internals-Git-Objects."),
        ("ref_dolt_block", "[8] DoltHub. 2024a. Dolt Storage Engine: Block Store. Dolt Documentation. https://www.dolthub.com/docs/architecture/storage-engine/block-store/."),
        ("ref_dolt_prolly", "[9] DoltHub. 2024b. Dolt's Storage Engine: Prolly Trees and Commit Graph. https://www.dolthub.com/blog/2024-02-29-storage-engine/."),
        ("ref_persistent", "[10] James R. Driscoll, Neil Sarnak, Daniel D. Sleator, and Robert E. Tarjan. 1989. Making Data Structures Persistent. Journal of Computer and System Sciences, 38(1):86-124. doi:10.1016/0022-0000(89)90034-2."),
        ("ref_orpheus", "[11] Silu Huang, Liqi Xu, Jialin Liu, Aaron J. Elmore, and Aditya G. Parameswaran. 2017. OrpheusDB: Bolt-on Versioning for Relational Databases. Proceedings of the VLDB Endowment, 10(10):1130-1141. doi:10.14778/3115404.3115417."),
        ("ref_decibel", "[12] Michael A. Maddox et al. 2016. Decibel: The Relational Dataset Branching System. Proceedings of the VLDB Endowment, 9(9):624-635. doi:10.14778/2947618.2947619."),
        ("ref_merkle", "[13] Ralph C. Merkle. 1988. A Digital Signature Based on a Conventional Encryption Function. In Advances in Cryptology - CRYPTO '87, LNCS 293, pages 369-378. doi:10.1007/3-540-48184-2_32."),
        ("ref_lbfs", "[14] Athicha Muthitacharoen, Benjie Chen, and David Mazieres. 2001. A Low-Bandwidth Network File System. Proceedings of the 18th ACM Symposium on Operating Systems Principles, pages 174-187. doi:10.1145/502034.502052."),
        ("ref_forkbase", "[15] Sheng Wang et al. 2018. ForkBase: An Efficient Storage Engine for Blockchain and Forkable Applications. Proceedings of the VLDB Endowment, 11(10):1137-1150. doi:10.14778/3231751.3231762."),
    ]
    for bookmark_id, (bookmark, reference) in enumerate(references, start=1):
        paragraph = doc.add_paragraph(style="IEEE Reference")
        paragraph.add_run(reference)
        add_bookmark(paragraph, bookmark, bookmark_id)

    add_section(doc, "Appendix A  Reproducibility and Data Availability")
    add_body(
        doc,
        "The repository contains the benchmark harness, evidence audit, trie-sensitivity harness, "
        "and paper generator. Source code, artifacts, and reproduction instructions are available "
        "at https://github.com/atharvasheersh/Chronos. The final evidence bundle is "
        "evidence/paper-final-20260824; "
        "the sensitivity bundle is evidence/trie-sensitivity-20260824. Results in this "
        "paper are generated programmatically from their summary CSV files and verified against "
        "raw results and manifests; no table or graph uses an earlier validation export.",
        indent=False,
    )

    # A terminal continuous section balances the final two-column reference page.
    balance_section = doc.add_section(WD_SECTION.CONTINUOUS)
    balance_section.page_width = Mm(210)
    balance_section.page_height = Mm(297)
    balance_section.top_margin = Inches(1.0)
    balance_section.bottom_margin = Inches(1.0)
    balance_section.left_margin = Inches(1.0)
    balance_section.right_margin = Inches(1.0)
    set_columns(balance_section, 1)

    for section in doc.sections:
        section.header.paragraphs[0].text = ""
        section.footer.paragraphs[0].text = ""
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        elif paragraph.style.name == "Caption":
            paragraph.paragraph_format.keep_with_next = paragraph.text.startswith("TABLE")

    link_citations(
        doc,
        {
            "[1]": "ref_delta",
            "[2]": "ref_noms",
            "[3]": "ref_mst",
            "[4]": "ref_ipfs",
            "[5]": "ref_datahub",
            "[6]": "ref_versioning",
            "[7]": "ref_git",
            "[8]": "ref_dolt_block",
            "[9]": "ref_dolt_prolly",
            "[10]": "ref_persistent",
            "[11]": "ref_orpheus",
            "[12]": "ref_decibel",
            "[13]": "ref_merkle",
            "[14]": "ref_lbfs",
            "[15]": "ref_forkbase",
        },
    )
    normalize_heading_following_indents(doc)
    enforce_submission_typography(doc)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
