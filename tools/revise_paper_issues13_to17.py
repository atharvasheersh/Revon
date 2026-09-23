"""Refresh the preserved submission manuscript from the issues 13–17 run."""

from __future__ import annotations

import csv
import json
import shutil
import statistics
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from build_final_research_paper import (  # noqa: E402
    make_calibration_chart,
    make_diff_chart,
    make_operational_chart,
    make_sensitivity_chart,
)

EVIDENCE = ROOT / "evidence" / "paper-corrected-issues13-17-counterbalanced-final-20260923"
SENSITIVITY = ROOT / "evidence" / "trie-sensitivity-issues8-20260923" / "summary.csv"
PAPER = ROOT / "paper" / "Revon_Final_Research_Paper.docx"
COVER = ROOT / "output" / "docs" / "Revon_Cover_Letter_Computing.docx"
TEMP = ROOT / "tmp" / "paper_issues13_to17"
SCENARIOS = (
    "small-sparse",
    "medium-sparse",
    "medium-dense",
    "large-sparse",
    "large-application-key-local",
    "large-range-local",
    "large-repeated-key",
    "large-hash-route-local",
)
LABELS = (
    "Small sparse",
    "Medium sparse",
    "Medium dense",
    "Large sparse",
    "Application-key local",
    "Range local",
    "Repeated key",
    "Hash-route local",
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def metric(row: dict[str, str], name: str) -> float:
    return float(row[f"{name}_median"])


def main() -> None:
    manifest = json.loads((EVIDENCE / "manifest.json").read_text(encoding="utf-8"))
    audit = json.loads((EVIDENCE / "evidence_audit.json").read_text(encoding="utf-8"))
    raw = read_csv(EVIDENCE / "raw_results.csv")
    summary = read_csv(EVIDENCE / "summary.csv")
    assert manifest["schema_version"] == 4 and manifest["profile"] == "paper"
    assert audit["passed"] and not audit["issues"]
    assert audit["execution_counts"] == {
        "warmups": 120,
        "measured": 420,
        "primary_evaluation_measured": 280,
        "dolt_evaluation_measured": 56,
        "dolt_bulk_evaluation_measured": 56,
    }
    assert len(raw) == 540 and len(summary) == 60
    assert all(row["status"] == "ok" and row["correctness"] == "True" for row in raw)
    assert all(row["execution_order"] for row in raw)
    assert manifest["models"] == ["snapshot", "log", "revon-m", "revon-h", "dolt", "dolt-bulk"]

    evaluation = {
        (row["scenario"], row["model"]): row
        for row in summary
        if row["phase"] == "evaluation"
    }
    calibration = [row for row in summary if row["phase"] == "calibration"]
    by_key = lambda scenario, model: evaluation[(scenario, model)]
    ratios = lambda metric_name, base="Dolt": [
        metric(by_key(s, base), metric_name) / metric(by_key(s, "Revon-H"), metric_name)
        for s in SCENARIOS
    ]
    diff_ratio = ratios("diff_ms")
    commit_ratio = ratios("incremental_commit_ms")
    checkout_ratio = ratios("checkout_ms")
    storage_ratio = [
        metric(by_key(s, "Revon-H"), "storage_bytes")
        / metric(by_key(s, "Dolt"), "storage_bytes")
        for s in SCENARIOS
    ]
    ablation_ratio = [
        metric(by_key(s, "Revon-M (forced Merkle)"), "diff_ms")
        / metric(by_key(s, "Revon-H"), "diff_ms")
        for s in SCENARIOS
    ]
    sql_import = statistics.median(metric(by_key(s, "Dolt"), "initial_import_ms") for s in SCENARIOS)
    bulk_import = statistics.median(
        metric(by_key(s, "Dolt (bulk import)"), "initial_import_ms") for s in SCENARIOS
    )
    revon_import = statistics.median(metric(by_key(s, "Revon-H"), "initial_import_ms") for s in SCENARIOS)
    sql_bulk_ratio = sql_import / bulk_import
    logical_ratio = statistics.median(
        metric(by_key(s, "Revon-H"), "storage_bytes_per_logical_payload_byte")
        / metric(by_key(s, "Dolt"), "storage_bytes_per_logical_payload_byte")
        for s in SCENARIOS
    )
    payload_residual = statistics.median(
        metric(by_key(s, "Revon-H"), "storage_bytes_minus_logical_payload_bytes")
        / metric(by_key(s, "Dolt"), "storage_bytes_minus_logical_payload_bytes")
        for s in SCENARIOS
    )
    compacted = {}
    for model in ("Revon-M (forced Merkle)", "Revon-H", "Dolt", "Dolt (bulk import)"):
        values = [
            int(row["storage_bytes_after_compaction"])
            for row in raw
            if row["phase"] == "evaluation"
            and row["trial_kind"] == "measured"
            and row["trial"] == "1"
            and row["model"] == model
            and row["storage_bytes_after_compaction"]
        ]
        compacted[model] = statistics.mean(values) / (1024 * 1024)
        assert len(values) == len(SCENARIOS)

    labels = dict(zip(SCENARIOS, LABELS))
    doc = Document(PAPER)

    def replace(prefix: str | tuple[str, ...], text: str) -> None:
        prefixes = (prefix,) if isinstance(prefix, str) else prefix
        matches = [p for p in doc.paragraphs if p.text.startswith(prefixes)]
        if not matches:
            matches = [p for p in doc.paragraphs if p.text == text]
        assert len(matches) == 1, (prefix, len(matches))
        p = matches[0]
        assert p.runs
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""

    def span(values: list[float]) -> str:
        return f"{min(values):.2f} to {max(values):.2f}"

    replace(
        "Versioning structured datasets requires",
        "Versioning structured datasets requires durable history, efficient sparse updates, historical reconstruction, and comparison without copying every record. Revon is a Python prototype built around an immutable fixed-depth Merkle hash trie. A commit rewrites affected leaf buckets and ancestor paths while reusing unchanged subtrees by content hash. Revon-H adds addressed changesets and selects log aggregation or hash-pruned tree differencing with a separately calibrated threshold. We evaluate Snapshot, Log-only, forced-Merkle Revon-M, Revon-H, and two Dolt 2.3.1 initial-import workflows on deterministic workloads up to 100,000 rows. The corrected run contains 540 executions: 120 warm-ups and 420 measured runs, including 280 primary five-system evaluation runs and 56 additional Dolt bulk-import runs. Across eight evaluation workloads, Dolt SQL/Revon-H median-latency ratios ranged from "
        + span(diff_ratio)
        + f" for diff and {span(commit_ratio)} for incremental commit. Dolt's bulk-import workflow is reported separately from SQL INSERT. Storage is reported as repository footprint and normalized per logical payload byte, not as pure metadata. These workflow measurements support a workload-dependent hybrid design, not general superiority over a production SQL database.",
    )
    replace(
        "A five-variant benchmark with shared",
        "A five-system benchmark with shared deterministic workloads, common canonical diff and checkout outputs, two separately reported Dolt initial-import workflows, balanced randomized execution order, external process-tree RSS sampling, raw evidence, and an internal evidence-integrity audit.",
    )
    replace(
        "Five implementations receive identical logical states",
        "Five comparison systems receive identical logical states and mutation batches: a complete Snapshot per version, an initial state plus operation-log baseline, forced-Merkle Revon-M, adaptive Revon-H, and Dolt 2.3.1 using a keyed SQL table and native JSON diff. Dolt initial import is tested two ways: batched SQL INSERT and CSV loaded with `dolt table import -r`; CSV serialization is included in the latter workflow's initial-import time. The import variants share the same Dolt engine for later operations, and their import times are kept distinct.",
    )
    replace(
        "Initial import measures the first durable state",
        "Initial import measures the first durable state; incremental commit is the within-trial median across ten commits; diff and checkout both materialize common outputs inside their timed regions. Diff is sorted changed keys with old/new SHA-256 value hashes encoded as canonical JSON. Checkout is the complete sorted key/value state encoded as canonical UTF-8 JSON. Dolt JSON parsing and normalization are included. Correctness compares the full state and hash contract with the workload oracle. Dolt SQL-path commit timing includes SQL mutation, add, and commit; commit hashes are retained for historical checkout and no version tags are created. Timed operations run without tracemalloc. An external parent process samples process-tree RSS every 10 ms for each isolated trial, including Dolt children; sampling is outside timed operations. Work-examined units differ by system and are not normalized.",
    )
    replace(
        "Each configuration has two warm-ups",
        f"Each configuration has two warm-ups and seven measured runs. Across calibration and evaluation, this gives 120 warm-ups and 420 measured runs. The primary five-system evaluation contains 280 measured runs, including 56 SQL-path Dolt runs; a further 56 measured evaluation runs use Dolt bulk import. Model order is randomized once per scenario and trial-kind, then rotated across trials. The first six measured trials balance all six variants across all six positions; order is recorded in raw_results.csv. Each system/trial runs in a clean worker process. One deterministic workload is reused per scenario; the internal audit regenerates it from seed {manifest['base_seed']} and checks its digest. Run {manifest['run_id']} was measured on {manifest['platform']} with Python {manifest['python_version']}. The audit recomputes the trial matrix and summary statistics; it does not independently reproduce timings.",
    )
    replace(
        "The 4,096-operation boundary is",
        f"The {manifest['hybrid_threshold_operations']:,}-operation boundary is a frozen policy for this machine, not a universal crossover. All 540 executions completed successfully and passed the harness's state and diff checks. The internal audit retained all {audit['outlier_candidates']} Tukey outlier candidates; no observation was manually deleted.",
    )

    replace(
        "Revon-H selected log for four sparse or hot workloads",
        f"Revon-H's Dolt/Revon-H median diff ratios ranged from {span(diff_ratio)} across the eight workloads; values above one favor Revon-H. The comparison uses separate medians without confidence intervals. Both systems materialize the same sorted key/hash JSON result. Dolt's CLI process and SQL costs remain inside this workflow comparison. The lexical-key workload is explicitly named application-key local; the other locality patterns are range local, repeated key, and hash-route local.",
    )
    replace(
        "For log-selected workloads, Revon-H's median diff",
        f"Across the eight workloads, forced-Merkle Revon-M/Revon-H median diff ratios ranged from {span(ablation_ratio)}. The results and selected path for every scenario are in Table 3 and the audited CSV. The within-system ablation shares the durable trie and therefore provides more direct evidence about adaptive selection than cross-system ratios.",
    )
    replace(
        "The Dolt/Revon-H median-latency ratio ranged from",
        f"Across the eight scenarios, Dolt SQL/Revon-H median-latency ratios ranged from {span(commit_ratio)} for incremental commit and {span(checkout_ratio)} for historical checkout. Across the same eight scenario medians, initial import took {sql_import:.2f} ms for the batched SQL INSERT workflow and {bulk_import:.2f} ms for Dolt bulk import; the SQL/bulk ratio was {sql_bulk_ratio:.2f} (above one means bulk import was faster). Revon-H's corresponding median initial import was {revon_import:.2f} ms. Per-scenario initial-import results are in the audited summary CSV. Bulk import includes creating the CSV input, `dolt table import -r`, and committing the table.",
    )
    replace(
        "Revon-H paid a consistent storage cost",
        f"Repository directory size is a whole-workflow footprint, not normalized algorithm storage. Across the eight scenarios, Revon-H/Dolt repository-footprint ratios ranged from {span(storage_ratio)}. Median Revon-H/Dolt ratios after dividing each directory by final UTF-8 logical key/value payload bytes were {logical_ratio:.2f}; the corresponding residual ratio was {payload_residual:.2f}. The residual is not pure metadata: it includes serialization, indexes, metadata, and compression effects. Post-compaction mean footprints for one measured sample per scenario were Revon-M {compacted['Revon-M (forced Merkle)']:.2f} MiB, Revon-H {compacted['Revon-H']:.2f} MiB, Dolt SQL {compacted['Dolt']:.2f} MiB, and Dolt bulk import {compacted['Dolt (bulk import)']:.2f} MiB. Dolt was compacted offline with `dolt gc` and Revon SQLite repositories with `VACUUM`; compaction time is outside latency and RSS. These are sampled outcomes, not a pure metadata comparison or universal steady-state result.",
    )
    replace(
        "RQ1 shows scenario-dependent workflow latency",
        "RQ1 shows scenario-dependent workflow latency and storage trade-offs across spread updates and three deliberately distinct locality patterns. Lexically adjacent application keys are not treated as generic hot-key locality. Range-local updates, repeated-key updates, and keys sharing Revon's hash-route prefix are separate scenarios. Results describe these synthetic workloads and do not imply range-query performance.",
    )
    replace(
        "Synthetic scope:",
        "Synthetic scope: deterministic key/value updates reach 100,000 rows and include application-key, sorted-range, repeated-key, and Revon hash-route locality; they do not represent every real schema, key skew, payload size, history length, or read-query pattern.",
    )
    replace(
        "Memory result: Across the five primary-evaluation scenarios",
        "Memory result: Across eight primary-evaluation scenarios, median sampled whole-worker process-tree RSS was collected for Snapshot, Log-only, Revon-M, Revon-H, and SQL-path Dolt. The values include common workload/oracle state and adapter processes, so they are not product-only memory. The parent sampled every 10 ms outside timed operations; peaks shorter than 10 ms may be missed, and the monitor can create small system-level scheduling pressure. Dolt bulk-import rows are not pooled with SQL-path Dolt for this comparison.",
    )
    replace(
        "Revon demonstrates a content-addressed versioning path",
        f"Revon demonstrates a content-addressed versioning path for structured key/value data: canonical immutable objects, incremental Merkle trie updates, atomic SQLite persistence, historical checkout, and adaptive differencing. Across eight scenarios, 280 primary measured evaluations compare five systems; 56 additional measured rows test Dolt's CSV bulk-import workflow. The paper reports separate SQL and bulk import times, scenario-level latency, normalized workflow storage, post-compaction samples, and external process-tree RSS. The within-system ablation provides more direct evidence about hybrid log/Merkle selection than cross-system ratios. These measurements are a reproducible prototype comparison, not a claim of production-system replacement.",
    )
    replace(
        "The repository contains the benchmark harness",
        f"The repository contains the benchmark harness, evidence audit, trie-sensitivity harness, and manuscript-formatting and validation tools. Source code, artifacts, and reproduction instructions are available at https://github.com/atharvasheersh/Revon. The corrected paper-profile evidence bundle is evidence/{EVIDENCE.name}; the trie-sensitivity bundle is evidence/trie-sensitivity-issues8-20260923. Tables and figures use the audited summary CSV; raw results record execution order and both Dolt initial-import workflows.",
    )

    def set_cell(cell, value: str) -> None:
        paragraph = cell.paragraphs[0]
        if paragraph.runs:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(value)

    workload = doc.tables[1]
    workload_specs = [
        ("Small sparse", "1,000", "10", "10", "spread"),
        ("Medium sparse", "10,000", "10", "10", "spread"),
        ("Medium dense", "10,000", "10", "1,000", "spread"),
        ("Large sparse", "100,000", "10", "100", "spread"),
        ("Application-key local", "100,000", "10", "100", "lexically adjacent keys"),
        ("Range local", "100,000", "10", "100", "contiguous sorted-key intervals"),
        ("Repeated key", "100,000", "10", "100", "repeated small key set"),
        ("Hash-route local", "100,000", "10", "100", "shared 9-bit SHA-256 prefix"),
    ]
    while len(workload.rows) < len(workload_specs) + 1:
        workload.add_row()
    for idx, values in enumerate(workload_specs, start=1):
        for col, value in enumerate(values):
            set_cell(workload.cell(idx, col), value)

    table = doc.tables[2]
    while len(table.rows) < len(SCENARIOS) + 1:
        table.add_row()
    set_cell(table.cell(0, 4), "SQL Dolt / H")
    for idx, (scenario, label) in enumerate(zip(SCENARIOS, LABELS), start=1):
        h, d = by_key(scenario, "Revon-H"), by_key(scenario, "Dolt")
        values = (
            label,
            h["strategy_selected"],
            f"{metric(h, 'diff_ms'):.3f}",
            f"{metric(d, 'diff_ms'):.2f}",
            f"{metric(d, 'diff_ms') / metric(h, 'diff_ms'):.1f}x",
        )
        for col, value in enumerate(values):
            set_cell(table.cell(idx, col), value)

    for p in doc.paragraphs:
        if p.text.startswith("Figure 4:"):
            p.runs[0].text = "Figure 4: Diff latency across eight scenarios, including separate application-key, range, repeated-key, and hash-route update locality; lower is better."
            for r in p.runs[1:]:
                r.text = ""
        elif p.text.startswith("Table 3:"):
            p.runs[0].text = "Table 3: Revon-H versus SQL-path Dolt median diff latency; bulk import is an initial-load variant and shares the same Dolt later-operation workflow."
            for r in p.runs[1:]:
                r.text = ""

    figure_caption_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith("Figure 4:"))
    table_caption_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith("Table 3:"))
    for p in doc.paragraphs[figure_caption_index + 1 : table_caption_index]:
        for page_break in p._p.xpath('.//w:br[@w:type="page"]'):
            page_break.getparent().remove(page_break)
    references_heading = next(p for p in doc.paragraphs if p.text.strip() == "References")
    references_heading.paragraph_format.page_break_before = True

    TEMP.mkdir(parents=True, exist_ok=True)
    assets = {
        "Figure 3:": ("calibration.png", lambda path: make_calibration_chart(path, calibration)),
        "Figure 4:": ("diff.png", lambda path: make_diff_chart(path, evaluation)),
        "Figure 5:": ("operational.png", lambda path: make_operational_chart(path, evaluation)),
        "Figure 6:": ("sensitivity.png", lambda path: make_sensitivity_chart(path, read_csv(SENSITIVITY))),
    }
    for caption, (filename, render) in assets.items():
        caption_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith(caption))
        image_paragraph = next(
            p for p in reversed(doc.paragraphs[:caption_index]) if p._p.xpath(".//a:blip")
        )
        blips = image_paragraph._p.xpath(".//a:blip")
        assert len(blips) == 1
        rid = blips[0].get(qn("r:embed"))
        path = TEMP / filename
        render(path)
        doc.part.related_parts[rid]._blob = path.read_bytes()

    backup = TEMP / f"before_{PAPER.name}"
    shutil.copy2(PAPER, backup)
    updated = TEMP / PAPER.name
    doc.save(updated)
    updated.replace(PAPER)
    print(PAPER)

    if COVER.exists():
        cover = Document(COVER)
        paragraphs = [p for p in cover.paragraphs if p.text.startswith("We evaluate Revon against snapshot")]
        assert len(paragraphs) == 1 and paragraphs[0].runs
        paragraphs[0].runs[0].text = (
            f"We evaluate Revon against snapshot and log-only baselines, a forced-Merkle ablation, and Dolt 2.3.1. "
            f"The corrected study contains 540 executions: 120 warm-ups and 420 measured runs, including 280 primary five-system evaluations and 56 additional Dolt bulk-import runs. "
            f"Across eight workloads, Dolt's SQL INSERT and `dolt table import -r` initial-import paths are reported separately. "
            f"The manuscript distinguishes total repository footprint from normalized bytes per logical payload and describes the tested locality patterns without claiming general superiority over a production SQL system."
        )
        for run in paragraphs[0].runs[1:]:
            run.text = ""
        shutil.copy2(COVER, TEMP / f"before_{COVER.name}")
        updated_cover = TEMP / COVER.name
        cover.save(updated_cover)
        updated_cover.replace(COVER)
        print(COVER)


if __name__ == "__main__":
    main()
