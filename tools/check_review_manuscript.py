"""Validate the expanded current paper while retaining the historical gates."""
from pathlib import Path
import re,json,hashlib,csv,statistics,zipfile
from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
MODEL=ROOT/'paper/Revon_Review_Publication_Model_20261016.json'
REPOSITORY='https://github.com/atharvasheersh/Revon'
ARCHIVE=REPOSITORY+'/blob/main/paper/Revon_Review_Revision_Reproducibility_20261016.zip'

def require(test,message):
    if not test:raise ValueError(message)

def check_doc(path,historical_model):
    spec=json.loads(MODEL.read_text(encoding='utf8'));doc=Document(path)
    paragraphs=[p.text for p in doc.paragraphs if p.text.strip()]
    require(paragraphs[0].casefold()==spec['title'].casefold(),'Springer title changed')
    require(len(doc.tables)==spec['tables'],'Springer table count changed')
    mapping={int(k):v for k,v in spec['citation_mapping'].items()}
    def remap(text):return re.sub(r'\[([\d, ]+)\]',lambda m:'['+', '.join(str(mapping[int(n)]) for n in re.findall(r'\d+',m[1]))+']',text)
    for i,(observed,expected) in enumerate(zip(doc.tables,spec['tables_cells']),1):
        require([[c.text for c in row.cells] for row in observed.rows]==expected,f'Springer table {i} differs from reviewed model')
    old_tables=[b for b in historical_model if b['kind']=='table']
    for expected in old_tables:
        number=int(spec['historical_table_numbers'][str(expected['number'])])
        cells=[[remap(c) for c in row] for row in [expected['headers']]+expected['rows']]
        require([[c.text for c in row.cells] for row in doc.tables[number-1].rows]==cells,f'historical table {expected["number"]} changed beyond citation remapping')
    require(len(doc.inline_shapes)==spec['figures'],'Springer figure count changed')
    for i,shape in enumerate(doc.inline_shapes,1):
        blob=doc.part.related_parts[shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))].blob
        require(hashlib.sha256(blob).hexdigest()==spec['figure_sha256'][str(i)],f'Springer figure {i} changed')
        if i<=7:
            with zipfile.ZipFile(ROOT/'paper/Revon_Linux_Reproducibility_Supplement.zip') as archive:
                require(blob==archive.read(f'figures/Fig{i}.png'),f'historical figure {i} changed')
    captions=[int(m[1]) for p in paragraphs if (m:=re.match(r'^Fig\.?\s+(\d+)\.',p))]
    require(captions==list(range(1,spec['figures']+1)),'Springer figure numbering changed')
    refs=[p for p in paragraphs if re.match(r'^\[\d+\]',p)]
    require(refs==spec['reference_paragraphs'],'Springer reference identity changed')
    numbers=[int(re.match(r'^\[(\d+)\]',p)[1]) for p in refs]
    require(numbers==list(range(1,spec['references']+1)),'Springer references changed')
    body='\n'.join(p for p in paragraphs if not re.match(r'^\[\d+\]',p))+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    cited={int(n) for group in re.findall(r'\[([\d, ]+)\]',body) for n in re.findall(r'\d+',group)}
    require(cited==set(numbers),'Springer uncited or undefined reference')
    for marker in ['516','72','93','3,402','2,646','32,768','16,384','65,536','poornima.n@vit.ac.in']:
        # Table totals may omit thousands separators while conclusion retains 3,402.
        require(marker in body or marker.replace(',','') in body,f'Springer missing evidence marker: {marker}')
    require('Online Resource 1' not in body,'Springer still claims journal supplementary hosting')
    availability=next(p for p in paragraphs if p.startswith('Data and code availability:'))
    start=next(i for i,p in enumerate(paragraphs) if p.startswith('APPENDIX A:'))
    end=next(i for i in range(start+1,len(paragraphs)) if paragraphs[i].startswith('APPENDIX B:'))
    # The methods appendix may legitimately discuss commits as a workload unit.
    access=paragraphs[start+1:end]
    require(not re.search(r'\bat commit\b|\bpinned at commit\b',availability+'\n'+'\n'.join(access)),'Springer availability still pins a repository commit')
    targets={rel.target_ref for rel in doc.part.rels.values() if rel.reltype==RT.HYPERLINK}
    require(REPOSITORY in targets and ARCHIVE in targets,'Springer current GitHub links missing')
    require(not any(re.search(r'github\.com/atharvasheersh/Revon/(?:tree|blob)/[0-9a-f]{40}(?:/|$)',t) for t in targets),'Springer GitHub link still pins a commit')
    require(paragraphs==spec['body_paragraphs'],'Springer paragraph differs from reviewed revision')
    require(len(next(p for p in paragraphs if p.startswith('Revon is a Python prototype')).split())<250,'Springer abstract too long')
    return {'tables':spec['tables'],'figures':spec['figures'],'references':spec['references'],'historical_tables_preserved':11},availability,access

def check_numeric_tables():
    spec=json.loads(MODEL.read_text(encoding='utf8'));analysis_path=ROOT/spec['campaign']/'analysis.json'
    require(hashlib.sha256(analysis_path.read_bytes()).hexdigest()==spec['analysis_sha256'],'review publication analysis changed')
    analysis=json.loads(analysis_path.read_text(encoding='utf8'))
    require(analysis['audit']['passed'] and analysis['audit']['completed_cases']==93,'review campaign incomplete')
    # Re-derive every numeric public/scale table row independently of the Word file.
    num=lambda x:f'{x:.3f}' if abs(x)<100 else f'{x:,.2f}'
    public=[];scale=[]
    for r in analysis['public_and_scale']:
        if r['dataset'] in ['co2-observed','wine-generated']:
            public.append([r['dataset'].split('-')[0]+'/'+r['model'],str(r['repositories']),num(r['diff_median_ms']),f"{num(r['diff_min_ms'])}-{num(r['diff_max_ms'])}",num(r['import_median_s']),num(r['storage_median_mib'])])
        else:scale.append([f"{int(r['dataset']):,}/{r['model']}"]+[num(r[k]) for k in ['import_median_s','commit_median_ms','diff_median_ms','checkout_median_s','rss_median_gib']])
    for key,rows in [('public',public),('scale',scale)]:
        n=spec['new_table_numbers'][key];require(spec['tables_cells'][n-1][1:]==rows,f'review {key} table differs from audited analysis')
    selector=list(csv.DictReader((analysis_path.parent/'selector.csv').open(encoding='utf8')))
    expected=[]
    for locality in ['spread','repeated-key']:
        for operations in [512,4096,32768]:
            selected=[r for r in selector if r['suite']=='validation' and r['locality']==locality and int(r['operations'])==operations]
            require(len(selected)==3,'held-out history count changed')
            expected.append([locality,f'{operations:,}']+[num(statistics.median(float(r[k]) for r in selected)) for k in ['log_ms','merkle_ms','hybrid4096_ms','calibrated_ms']])
    n=spec['new_table_numbers']['selector']
    require(spec['tables_cells'][n-1][1:]==expected,'review selector table differs from audited history medians')
    return {'new_numeric_tables_match_audited_analysis':True}
