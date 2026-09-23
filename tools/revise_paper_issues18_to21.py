"""Update the final manuscript with issues 18–21 evidence and limits."""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "evidence" / "paper-corrected-issues13-17-counterbalanced-final-20260923"
WINDOWS = ROOT / "evidence" / "paper-issues18-20-robustness-20260923"
LINUX = ROOT / "evidence" / "paper-issues18-20-linux-wsl-20260923"
PAPER = ROOT / "paper" / "Revon_Final_Research_Paper.docx"
TEMP = ROOT / "tmp" / "paper_issues18_to21"
SCENARIOS = (
    "small-sparse",
    "medium-sparse",
    "medium-dense",
    "large-sparse",
    "large-application-key-local",
    "large-range-local",
    "large-repeated-key",
    "large-hash-route-local",
)
LABELS = (
    "Small sparse",
    "Medium sparse",
    "Medium dense",
    "Large sparse",
    "Application-key local",
    "Range local",
    "Repeated key",
    "Hash-route local",
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as stream:
        return list(csv.DictReader(stream))


def main() -> None:
    manifest = json.loads((EVIDENCE / "manifest.json").read_text(encoding="utf-8"))
    audit = json.loads((EVIDENCE / "evidence_audit.json").read_text(encoding="utf-8"))
    primary = read_csv(EVIDENCE / "summary.csv")
    uncertainty = read_csv(EVIDENCE / "paired_ratio_uncertainty.csv")
    win_manifest = json.loads((WINDOWS / "manifest.json").read_text(encoding="utf-8"))
    win_audit = json.loads((WINDOWS / "audit.json").read_text(encoding="utf-8"))
    linux_manifest = json.loads((LINUX / "manifest.json").read_text(encoding="utf-8"))
    linux_audit = json.loads((LINUX / "audit.json").read_text(encoding="utf-8"))
    win_factors = read_csv(WINDOWS / "factor_bootstrap.csv")
    linux_factors = read_csv(LINUX / "factor_bootstrap.csv")
    assert audit["passed"] and not audit["issues"]
    assert win_audit["passed"] and linux_audit["passed"]
    assert len(uncertainty) == 48
    assert win_audit["rows"] == 405 and win_audit["measured_rows"] == 315
    assert linux_audit["rows"] == 180 and linux_audit["measured_rows"] == 135
    assert len(win_factors) == len(linux_factors) == 5

    primary_summary = {
        (row["scenario"], row["model"]): row
        for row in primary
        if row["phase"] == "evaluation"
    }
    paired = {
        (row["scenario"], row["numerator_model"], row["metric"]): row
        for row in uncertainty
    }
    dolt_diff_rows = [
        paired[(scenario, "Dolt", "diff_ms")] for scenario in SCENARIOS
    ]
    paired_ratio_min = min(float(row["median_paired_ratio"]) for row in dolt_diff_rows)
    paired_ratio_max = max(float(row["median_paired_ratio"]) for row in dolt_diff_rows)
    merkle_diff_rows = [
        paired[(scenario, "Revon-M (forced Merkle)", "diff_ms")]
        for scenario in SCENARIOS
    ]
    parity_crossing = sum(
        float(row["ci95_low"]) <= 1 <= float(row["ci95_high"])
        for row in merkle_diff_rows
    )
    assert parity_crossing == 3
    assert all(float(row["ci95_low"]) > 1 for row in dolt_diff_rows)

    win_by_factor = {row["workload_factor"]: row for row in win_factors}
    linux_by_factor = {row["workload_factor"]: row for row in linux_factors}
    factor_order = ("baseline", "payload-128", "history-50", "range-local", "repeated-key")
    assert all(factor in win_by_factor and factor in linux_by_factor for factor in factor_order)

    doc = Document(PAPER)
    if (
        any(p.text.startswith("The primary matrix has two warm-ups") for p in doc.paragraphs)
        and doc.tables[2].cell(0, 4).text == "Dolt/H ratio [95% CI]"
    ):
        print(f"Already revised: {PAPER}")
        return

    def replace(prefix: str, text: str) -> None:
        matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
        assert len(matches) == 1, (prefix, len(matches))
        paragraph = matches[0]
        assert paragraph.runs
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""

    def ratio_text(row: dict[str, str]) -> str:
        return f"{float(row['median_paired_ratio']):.1f} ({float(row['ci95_low']):.1f}-{float(row['ci95_high']):.1f})"

    def factor_ci(row: dict[str, str]) -> str:
        return (
            f"{float(row['median_paired_ratio_rev_m_over_rev_h']):.2f} "
            f"[{float(row['ci95_low']):.2f}, {float(row['ci95_high']):.2f}]"
        )

    replace(
        "Versioning structured datasets requires durable history",
        f"Versioning structured datasets requires durable history, efficient sparse updates, historical reconstruction, and comparison without copying every record. Revon is a Python prototype built around an immutable fixed-depth Merkle hash trie; Revon-H adds addressed changesets and adaptively selects log aggregation or hash-pruned differencing. We evaluate Snapshot, Log-only, forced-Merkle Revon-M, Revon-H, and Dolt 2.3.1 on deterministic workloads up to 100,000 rows. The primary run contains 540 executions. Across eight workloads, median paired Dolt/Revon-H diff ratios range from {paired_ratio_min:.2f} to {paired_ratio_max:.2f}; all seven-trial 95% bootstrap intervals favor Revon-H latency in this tested workflow. A separate three-seed study across payload sizes, history lengths, and update patterns finds the 50-commit result inconclusive. These measurements compare the declared workflows on one Windows host and WSL2 on the same physical computer. They support a workload-dependent prototype design, not general or production-level superiority.",
    )

    primary_methods = next(p for p in doc.paragraphs if p.text.startswith("Each configuration has two warm-ups"))
    primary_methods.runs[0].text = (
        f"The primary matrix has two warm-ups and seven measured trials per configuration: 120 warm-ups and 420 measured runs across calibration and evaluation. The primary five-system evaluation contains 280 measured runs, including 56 SQL-path Dolt runs; 56 additional measured runs use Dolt bulk import. Model order is counterbalanced and recorded. Each system/trial runs in a clean worker. The primary run reuses one deterministic workload per scenario, regenerated and digest-checked by the internal audit. Run {manifest['run_id']} used {manifest['platform']} and Python {manifest['python_version']}. For each scenario and comparison, latency ratios are formed from matching trial numbers; a 20,000-resample percentile bootstrap of the seven paired ratios supplies a 95% interval. These intervals describe run-to-run variation conditional on the fixed workload and host, not variation across datasets or machines. Intervals that include parity are classified as inconclusive; a 5% practical margin is used as an analyst-selected label, not as an external standard."
    )
    for run in primary_methods.runs[1:]:
        run.text = ""

    replace(
        "The 4,096-operation boundary is",
        f"The {manifest['hybrid_threshold_operations']:,}-operation boundary is a frozen policy for this machine, not a universal crossover. All 540 primary executions completed successfully and passed state and diff checks. The audit retained all {audit['outlier_candidates']} Tukey outlier candidates; no observation was deleted. A separate sensitivity run varies seed, payload size, history length, and update pattern; it is kept separate from the primary tables.",
    )

    replace(
        "Table 3: Revon-H versus SQL-path Dolt median diff latency",
        "Table 3: Revon-H versus SQL-path Dolt diff latency; the final column reports the median paired trial ratio and percentile-bootstrap 95% interval.",
    )
    replace(
        "Revon-H's Dolt/Revon-H median diff ratios ranged",
        "Across the eight scenarios, every paired Dolt/Revon-H diff ratio is above one and its 95% interval excludes parity, indicating lower Revon-H latency for this tested workflow. Table 3 reports the median of the seven same-trial ratios and its bootstrap interval. This remains an in-process Revon API versus Dolt CLI workflow comparison: Dolt's CLI and SQL costs are included, so the ratios are not engine-only speedups. The lexical-key case is application-key local; range-local, repeated-key, and hash-route-local updates remain distinct workloads.",
    )

    replace(
        "Across the eight workloads, forced-Merkle Revon-M/Revon-H median diff ratios ranged",
        f"In the primary run, {parity_crossing} of eight paired Revon-M/Revon-H diff intervals include parity and are inconclusive. A separate robustness matrix covers 10,000 rows, three seeds, 32- and 128-byte values, 10- and 50-commit histories, and spread, range-local, and repeated-key updates. Revon-M/Revon-H diff ratios above one mean Revon-M was slower. On Windows, baseline, 128-byte payload, range-local, and repeated-key ratios were {factor_ci(win_by_factor['baseline'])}, {factor_ci(win_by_factor['payload-128'])}, {factor_ci(win_by_factor['range-local'])}, and {factor_ci(win_by_factor['repeated-key'])}. The 50-commit interval crossed parity. WSL2 showed the same direction for those four factors and parity overlap for the 50-commit history. The paired ablation is more direct evidence about adaptive selection than cross-system ratios, but this three-seed, one-host sensitivity study does not establish general superiority.",
    )

    replace(
        "Across the eight scenarios, Dolt SQL/Revon-H median-latency ratios ranged",
        "Across all eight scenarios, paired Dolt/Revon-H intervals exclude parity for initial import, incremental commit, diff, and historical checkout. The full paired ratios and intervals are in paired_ratio_uncertainty.csv. Initial-import workflow medians across scenario medians were 1,814.83 ms for batched SQL INSERT, 1,701.25 ms for Dolt bulk import, and 1,131.24 ms for Revon-H. In the paired SQL/bulk-import analysis, four scenario intervals favor bulk import and four include parity; there is no general bulk-import advantage. These are workflow measurements; they do not establish engine-level or production-system superiority.",
    )

    replace(
        "RQ1 shows scenario-dependent workflow latency",
        "RQ1 shows scenario-dependent workflow latency and storage trade-offs across the declared synthetic update patterns. A separate three-seed sensitivity study found consistent Revon-H diff advantages over the forced-Merkle ablation for four tested factors on Windows and WSL2. The 50-commit case remained inconclusive on both. Lexically adjacent application keys, sorted-range updates, repeated-key updates, and hash-route locality are different update patterns. None of these experiments measures ordered range queries, concurrent reads or writes, or production throughput.",
    )
    replace(
        "RQ2 is supported more directly.",
        "RQ2 is supported more directly for the tested short and sparse histories: the paired Revon-M/Revon-H ablation shares the durable trie, while the added 50-commit case that crossed the frozen operation boundary showed no resolved difference. This does not establish a universal selector advantage.",
    )
    replace(
        "Synthetic scope:",
        "Synthetic scope: the primary study reaches 100,000 rows; the additional sensitivity study uses 10,000-row flat key-value tables, three seeds, 32- and 128-byte values, 10- and 50-commit histories, and three update patterns. It does not cover alternate schemas, public multi-table datasets, or one-million-row data.",
    )
    replace(
        "Single environment:",
        "Environment scope: the primary run used one Windows 11 host. The sensitivity run also used WSL2 Ubuntu on that same physical computer; this checks a second software environment but is not an independent-machine replication.",
    )
    replace(
        "Outlier policy:",
        "Uncertainty: all Tukey candidates were retained. The primary seven-trial ratio intervals are paired bootstrap intervals conditional on each fixed scenario. The additional seed-cluster bootstrap used only three seeds. Neither analysis establishes performance over a population of datasets or machines.",
    )
    replace(
        "To move beyond a linear prototype,",
        "To move beyond a linear prototype, Revon still needs named branches, merge commits, conflict reporting, schema-aware records, concurrent readers and writers, ordered range scans, and remote synchronization. Evaluation still needs public structured datasets, one-million-row workloads, and an independent Linux machine. WSL2 added process-tree CPU sampling and serial commit-operation throughput; its I/O byte counters were not available for every trial and included setup and correctness work, not only the timed database operation. A full cross-system resource study therefore remains open.",
    )
    replace(
        "Revon demonstrates a content-addressed versioning path",
        "Revon demonstrates a content-addressed versioning path for structured key/value data through immutable objects, incremental Merkle updates, SQLite persistence, historical checkout, and adaptive differencing. Paired intervals and a three-seed sensitivity study support workload-specific differences between Revon-H and the forced-Merkle ablation; the long-history result remains inconclusive. The primary benchmark and follow-up runs do not establish performance across real schemas, public datasets, independent hardware, concurrent workloads, or ordered range queries. These measurements describe a reproducible prototype, not a production-system replacement.",
    )
    replace(
        "The repository contains the benchmark harness",
        f"The repository contains the benchmark harness, evidence audit, paired-ratio analysis, multi-seed robustness study, and manuscript validation tools. The primary corrected paper-profile bundle is evidence/{EVIDENCE.name}. The primary paired ratio intervals are in evidence/{EVIDENCE.name}/paired_ratio_uncertainty.csv. The separate Windows and WSL2 sensitivity bundles are evidence/{WINDOWS.name} and evidence/{LINUX.name}; their raw rows, manifests, audits, summaries, and seed-cluster intervals are retained. The trie-sensitivity bundle is evidence/trie-sensitivity-issues8-20260923. These artifacts validate evidence consistency; they do not independently reproduce timings. Project repository: ",
    )
    repro = next(p for p in doc.paragraphs if p.text.startswith("The repository contains the benchmark harness"))
    repro.alignment = WD_ALIGN_PARAGRAPH.LEFT

    table = doc.tables[2]
    table.autofit = False
    for column, width in zip(table.columns, (1.56, 0.65, 0.95, 1.10, 1.45)):
        column.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, (1.56, 0.65, 0.95, 1.10, 1.45)):
            cell.width = Inches(width)
    set_cell = lambda cell, value: _set_cell(cell, value)
    set_cell(table.cell(0, 4), "Dolt/H ratio [95% CI]")
    for row_index, scenario in enumerate(SCENARIOS, start=1):
        row = paired[(scenario, "Dolt", "diff_ms")]
        set_cell(table.cell(row_index, 4), ratio_text(row))
        for run in table.cell(row_index, 4).paragraphs[0].runs:
            run.font.size = Pt(8)

    TEMP.mkdir(parents=True, exist_ok=True)
    backup = TEMP / f"before_{PAPER.name}"
    shutil.copy2(PAPER, backup)
    updated = TEMP / PAPER.name
    doc.save(updated)
    updated.replace(PAPER)
    print(PAPER)


def _set_cell(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


if __name__ == "__main__":
    main()
