from pathlib import Path
import re, io, json, hashlib, uuid, zipfile
from copy import deepcopy
from xml.sax.saxutils import escape
from collections import Counter
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph as WordParagraph
from fontTools.ttLib import TTFont
from reportlab.pdfbase import pdfmetrics, rl_codecs
from reportlab.pdfbase.pdfmetrics import EmbeddedType1Face, Font, Encoding
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Image, Table, TableStyle, KeepTogether
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
from reportlab.lib.colors import black
from PIL import Image as PILImage

ROOT=Path(__file__).resolve().parents[2]
TMP=Path(__file__).resolve().parent
OUT=ROOT/'output'/'new format'
SOURCE=ROOT/'paper'/'Revon_Research_Paper_Discover_Computing.docx'
REFERENCE=Path(r'C:\Users\admin\Downloads\Use for formatting, TFIDF Paper.pdf')
STEM='Revon_Research_Paper_TFIDF_Reference_Format'
L,R,W=102.875,104.625,404.5
FS,LEAD,SMALL,SLEAD=10.9589,12.9514,8.9664,10.9589
source=Document(SOURCE)
src_pars=source.paragraphs
source_hash=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
family='TeXGyreTermesX'

class SharedType1Face(EmbeddedType1Face):
    def addObjects(self, doc):
        if not hasattr(self, '_descriptor_cache'):
            self._descriptor_cache={}
        if id(doc) not in self._descriptor_cache:
            self._descriptor_cache[id(doc)]=super().addObjects(doc)
        return self._descriptor_cache[id(doc)]

def prepare_fonts():
    # Use the original full Type 1 programs corresponding to the reference's
    # TeXGyreTermesX fonts, not a Times substitute.
    cmap=list(pdfmetrics.getEncoding('WinAnsiEncoding').vector)
    cmap[129]='lslash'
    enc=Encoding('RevonWinAnsi',cmap)
    pdfmetrics.registerEncoding(enc)
    smallcaps=cmap.copy()
    for k in range(97,123):smallcaps[k]=chr(k)+'.sc'
    pdfmetrics.registerEncoding(Encoding('RevonSmallCaps',smallcaps))
    decoding={}
    for k in range(256):
        try: v=ord(bytes([k]).decode('cp1252'))
        except UnicodeDecodeError: v=None
        decoding[k]=v
    decoding[129]=322
    encoding={v:k for k,v in decoding.items() if v is not None}
    rl_codecs.RL_Codecs.add_dynamic_codec('RevonWinAnsi',encoding,decoding)
    rl_codecs.RL_Codecs.add_dynamic_codec('RevonSmallCaps',encoding,decoding)
    rl_codecs.RL_Codecs.register()
    for style,pfb in [('Regular','ztmr'),('Bold','ztmb'),('Italic','ztmri'),('BoldItalic','ztmbi')]:
        font=TTFont(TMP/f'TeXGyreTermesX-{style}.otf')
        up=font['head'].unitsPerEm
        def cv(x): return round(x*1000/up)
        afm=TMP/f'{style}.afm'
        name=font['name'].getDebugName(6)
        glyphs=font.getGlyphOrder()
        rows=['StartFontMetrics 2.0',f'FontName {name}',f'FullName {name}',
              'FamilyName TeXGyreTermesX',f'Weight {style}',
              f'ItalicAngle {font["post"].italicAngle}',f'Ascender {cv(font["hhea"].ascent)}',
              f'Descender {cv(font["hhea"].descent)}',
              f'FontBBox {cv(font["head"].xMin)} {cv(font["head"].yMin)} {cv(font["head"].xMax)} {cv(font["head"].yMax)}',
              f'CapHeight {cv(font["OS/2"].sCapHeight)}',f'StartCharMetrics {len(glyphs)}']
        for g in glyphs:
            rows.append(f'C -1 ; WX {cv(font["hmtx"][g][0])} ; N {g} ; B 0 0 0 0 ;')
        rows+=['EndCharMetrics','EndFontMetrics']
        afm.write_text('\n'.join(rows),encoding='ascii')
        face=SharedType1Face(str(afm),str(TMP/f'{pfb}.pfb'))
        pdfmetrics.registerTypeFace(face)
        pdfmetrics.registerFont(Font('TG'+style,face.name,'RevonWinAnsi'))
        pdfmetrics.registerFont(Font('TG'+style+'SC',face.name,'RevonSmallCaps'))
    pdfmetrics.registerFontFamily('TGRegular',normal='TGRegular',bold='TGBold',italic='TGItalic',boldItalic='TGBoldItalic')
    pdfmetrics.registerFontFamily('TGRegularSC',normal='TGRegularSC',bold='TGBoldSC',italic='TGItalicSC',boldItalic='TGBoldItalicSC')

prepare_fonts()

# Alphabetical, numbered bibliography, as in the supplied manuscript.
refs=[]
for p in src_pars:
    m=re.match(r'^\[(\d+)\]\s+(.*)',p.text)
    if m:
        old=int(m[1]); text=m[2]
        title=re.search(r'[“"](.+?)[”"]',text)
        author=text[:title.start()].rstrip(', ') if title else text.split(',')[0]
        refs.append((old,text,author))
def surname_key(v):
    first=v[2].split(',')[0].split(' and ')[0]
    first=re.sub(r'^(?:[A-Z]\.\s*)+','',first)
    return re.sub(r'[^a-z]','',first.lower())
refs.sort(key=surname_key)
mapping={old:i+1 for i,(old,_,_) in enumerate(refs)}

def remap(t):
    t=re.sub(r'\[([\d,\s]+)\]',lambda m:'['+', '.join(str(mapping.get(int(n),int(n))) for n in re.findall(r'\d+',m[1]))+']',t)
    t=re.sub(r'\[\d+\](?:,\s*\[\d+\])+',lambda m:'['+', '.join(re.findall(r'\d+',m[0]))+']',t)
    return t

def sc(t,size=FS,italic=False):
    # Use native small-cap glyphs in the original Type 1 font program.
    return f'<font name="TG{"Italic" if italic else "Regular"}SC" size="{size:.5f}">'+escape(t)+'</font>'

def linked(t):
    t=escape(t)
    t=re.sub(r'(https?://[^\s&lt;]+)',lambda m:'<link href="'+m[0].rstrip('.,')+'" color="blue">'+m[0]+'</link>',t)
    return t

def body_markup(t):
    t=linked(remap(t))
    # Numeric references retain their identities after sorting the bibliography.
    return re.sub(r'\[([\d, ]+)\]',lambda m:'['+', '.join('<link href="#ref'+n+'" color="blue">'+n+'</link>' for n in re.findall(r'\d+',m[1]))+']',t)

records=[]
table_index=0
for el in source.element.body:
    if el.tag==qn('w:p'):
        p=WordParagraph(el,source)
        idx=next((i for i,x in enumerate(src_pars) if x._p is el),None)
        if idx is not None and idx<8: continue
        if p._p.xpath('.//w:drawing'):
            b=p._p.xpath('.//a:blip')[0]
            image_bytes=source.part.related_parts[b.get(qn('r:embed'))].blob
            records.append({'kind':'image','data':image_bytes})
        elif p.text.strip(): records.append({'kind':'p','text':p.text,'style':p.style.name,'idx':idx})
    elif el.tag==qn('w:tbl'):
        table=WordTable(el,source)
        records.append({'kind':'table','cells':[[c.text for c in r.cells] for r in table.rows],'number':table_index+1})
        table_index+=1

appendix=[]
start=next(i for i,r in enumerate(records) if r.get('text','').startswith('APPENDIX A.'))
end=next(i for i in range(start+1,len(records)) if records[i].get('style')=='Heading 1')
appendix=records[start:end]
records=records[:start]+records[end:]
records=records[:next(i for i,r in enumerate(records) if r.get('text')=='REFERENCES')]

styles={
 'body':ParagraphStyle('Body',fontName='TGRegular',fontSize=FS,leading=LEAD,alignment=TA_JUSTIFY,firstLineIndent=11.955,spaceAfter=0,splitLongWords=True,uriWasteReduce=.2),
 'abstract':ParagraphStyle('Abstract',fontName='TGRegular',fontSize=SMALL,leading=SLEAD,alignment=TA_JUSTIFY,leftIndent=64.765,rightIndent=64.765,firstLineIndent=15.94,spaceAfter=19.925),
 'title':ParagraphStyle('Title',fontName='TGBold',fontSize=FS,leading=LEAD,alignment=TA_CENTER,spaceAfter=18),
 'author':ParagraphStyle('Authors',fontName='TGRegular',fontSize=FS,leading=LEAD,alignment=TA_CENTER,spaceAfter=11.5),
 'affil':ParagraphStyle('Affiliation',fontName='TGItalic',fontSize=7.9701,leading=9.9626,alignment=TA_CENTER,spaceAfter=9),
 'caption':ParagraphStyle('Caption',fontName='TGItalic',fontSize=SMALL,leading=SLEAD,alignment=TA_JUSTIFY,spaceBefore=5,spaceAfter=14),
 'tablabel':ParagraphStyle('TableLabel',fontName='TGItalic',fontSize=SMALL,leading=SLEAD,alignment=TA_CENTER,spaceBefore=12,spaceAfter=0),
 'tabcaption':ParagraphStyle('TableCaption',fontName='TGItalic',fontSize=SMALL,leading=SLEAD,alignment=TA_CENTER,spaceAfter=5),
 'ref':ParagraphStyle('Reference',fontName='TGRegular',fontSize=SMALL,leading=SLEAD,leftIndent=36,firstLineIndent=-36,alignment=TA_JUSTIFY,spaceAfter=1),
 'center':ParagraphStyle('Center',fontName='TGRegular',fontSize=FS,leading=LEAD,alignment=TA_CENTER,spaceBefore=16,spaceAfter=8),
}

kw=src_pars[7].text.replace('Keywords:','').strip()
def decorate(canvas,doc):
    canvas.saveState()
    canvas.setFont('TGRegular',9.9626)
    if doc.page==1:
        canvas.drawCentredString((L+612-R)/2,66.5,'1')
        canvas.setLineWidth(.45);canvas.line(L,110,L+161.4,110)
        st=ParagraphStyle('KW',fontName='TGRegular',fontSize=SMALL,leading=SLEAD,firstLineIndent=11.955)
        p=Paragraph('<i>Keywords and phrases:</i> '+escape(kw)+'.',st)
        _,h=p.wrap(W,100);p.drawOn(canvas,L,80)
    else:
        if doc.page%2:canvas.drawRightString(612-R,739.6,str(doc.page))
        else:canvas.drawString(L,739.6,str(doc.page))
    canvas.restoreState()

pdf=OUT/(STEM+'.pdf')
doc=BaseDocTemplate(str(pdf),pagesize=(612,792),leftMargin=L,rightMargin=R,
                    topMargin=65.0,bottomMargin=77.5,title=src_pars[0].text,
                    author=src_pars[1].text,pageCompression=1)
doc.addPageTemplates([
    PageTemplate(id='First',frames=[Frame(L,123,W,603.925,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=decorate,autoNextPageTemplate='Body'),
    PageTemplate(id='Body',frames=[Frame(L,77.5,W,650.175,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=decorate),
])
story=[Spacer(1,18),Paragraph(src_pars[0].text.upper(),styles['title']),
       Paragraph(sc('By Atharva Sheersh Pandey')+'<super><font size="7.6712">1,a</font></super>, '+sc('Adhyan Jain')+'<super><font size="7.6712">1,b</font></super>,<br/>'+sc('and Poornima Nedunchezhian')+'<super><font size="7.6712">1,c</font></super>',styles['author']),
       Paragraph('<super><font size="6.7746">1</font></super>'+sc('School of Computer Science and Engineering, VIT Vellore',7.9701,True),styles['affil']),
       Paragraph('<super>a</super><link href="mailto:atharva.sheersh2024@vitstudent.ac.in" color="blue">atharva.sheersh2024@vitstudent.ac.in</link>; <super>b</super><link href="mailto:adhyan.jain2024@vitstudent.ac.in" color="blue">adhyan.jain2024@vitstudent.ac.in</link><br/><super>c</super>'+sc('Corresponding author:',7.9701,True)+' <link href="mailto:poornima.n@vit.ac.in" color="blue">poornima.n@vit.ac.in</link>',styles['affil']),
       Paragraph(escape(src_pars[6].text.removeprefix('Abstract: ')),styles['abstract'])]

def make_table(rec):
    cells=rec['cells'];n=len(cells[0]);num=rec['number']
    widths={1:[91,144,169.5],2:[92,312.5],3:[145,62,105,92.5],4:[117,287.5],
            5:[28,55,55,194.5,72],6:[106,71,83,65,79.5],7:[34,75,75,220.5],
            8:[30,68,68,80,74,84.5],9:[54,66,78,65,76,65.5],10:[60,55,67,82,73,67.5],
            11:[86,57,67,62,73,59.5]}[num]
    size=7.9701 if n>=5 or num==1 else SMALL
    # Hashes wrap without losing any characters.
    cellstyle=ParagraphStyle('Cell',fontName='TGRegular',fontSize=size,leading=size+1.5,splitLongWords=True,alignment=TA_LEFT)
    rows=[]
    for ri,row in enumerate(cells):
        rows.append([Paragraph(body_markup(t),cellstyle) for t in row])
    t=Table(rows,colWidths=widths,repeatRows=1,hAlign='CENTER')
    t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),3),
                          ('RIGHTPADDING',(0,0),(-1,-1),3),('TOPPADDING',(0,0),(-1,-1),2),
                          ('BOTTOMPADDING',(0,0),(-1,-1),2),('LINEBELOW',(0,0),(-1,0),.5,black),
                          ('LINEBELOW',(0,-1),(-1,-1),.5,black)]))
    return t

main=0;sub=0;i=0
while i<len(records):
    r=records[i]
    if r['kind']=='image':
        im=PILImage.open(io.BytesIO(r['data']));h=W*im.height/im.width
        figure=Image(io.BytesIO(r['data']),width=W,height=h)
        pair=[Spacer(1,12),figure]
        if i+1<len(records) and records[i+1].get('style')=='Caption':
            text=records[i+1]['text'];m=re.match(r'Fig\. (\d+)\s+(.*)',text)
            pair.append(Paragraph(sc('Fig',SMALL,True)+' '+m[1]+'. '+body_markup(m[2]),styles['caption']));i+=1
        story.append(KeepTogether(pair))
    elif r['kind']=='table':
        story+=[make_table(r),Spacer(1,12)]
    elif r.get('style')=='Caption':
        m=re.match(r'Table (\d+)\s+(.*)',r['text'])
        if m:
            a=Paragraph(sc('Table',SMALL,True)+' '+m[1],styles['tablabel'])
            b=Paragraph(body_markup(m[2]),styles['tabcaption'])
            a.keepWithNext=True;b.keepWithNext=True;story+=[a,b]
    elif r.get('style','').startswith('Heading'):
        text=r['text']
        if text=='DECLARATIONS': i+=1;continue
        h2=r['style']=='Heading 2'
        if h2:sub+=1;label=f'{main}.{sub}. '+text+'.'
        elif re.match(r'^\d+ ',text):
            main=int(text.split()[0]);sub=0;label=re.sub(r'^(\d+) ',r'\1. ',text)+'.'
        else:label=text.capitalize()+'.'
        fmt='i' if h2 else 'b'
        inline=''
        if i+1<len(records) and records[i+1]['kind']=='p' and records[i+1]['style']=='Normal':
            inline='&nbsp;&nbsp; '+body_markup(records[i+1]['text']);i+=1
        st=ParagraphStyle('RunIn',parent=styles['body'],spaceBefore=13 if not h2 else 10,keepWithNext=not bool(inline))
        story.append(Paragraph(f'<{fmt}>'+escape(label)+f'</{fmt}>'+inline,st))
    else:
        text=r['text'];mark=body_markup(text)
        if r.get('idx') in range(135,141):
            lead,body=text.split(':',1);mark='<b>'+escape(lead)+'.</b>&nbsp;&nbsp; '+body_markup(body.strip())
        st=styles['body']
        if r.get('idx') in range(135,141): st=ParagraphStyle('Decl',parent=st,spaceBefore=10)
        story.append(Paragraph(mark,st))
    i+=1

story.append(Paragraph('REFERENCES',styles['center']))
for old,text,author in refs:
    # Keep each bibliographic fact and identifier; only move the year to the
    # author position and replace quotation marks with reference-style prose.
    yr=re.search(r'(?<!\d)(19|20)\d{2}(?!\d)',text)
    year=yr[0] if yr else None
    tail=text[len(author):].lstrip(', ')
    if year:tail=re.sub(r'(?<!\d)'+year+r'(?!\d)[,.]?\s*','',tail,count=1)
    tail=tail.replace('“','').replace('”','').replace('"','')
    html=f'<a name="ref{mapping[old]}"/>[{mapping[old]}]&nbsp;&nbsp; '+sc(author,SMALL)+(f' ({year}). ' if year else '. ')+linked(tail)
    story.append(Paragraph(html,styles['ref']))
story.append(Paragraph('APPENDIX A: REPRODUCIBILITY',styles['center']))
for r in appendix[1:]: story.append(Paragraph(body_markup(r['text']),styles['body']))
doc.build(story)

# Editable companion retains the original OOXML figures, links, and tables.
word=Document(SOURCE)
for section in word.sections:
    section.page_width=Pt(612);section.page_height=Pt(792)
    section.left_margin=Pt(L);section.right_margin=Pt(R)
    section.top_margin=Pt(65);section.bottom_margin=Pt(77.5)
    section.header_distance=Pt(40);section.footer_distance=Pt(62)
    section.different_first_page_header_footer=True
    cols=section._sectPr.find(qn('w:cols'))
    if cols is not None:cols.set(qn('w:num'),'1')
word.settings.odd_and_even_pages_header_footer=True
for name in ['Normal','Title','Heading 1','Heading 2','Caption','Default Paragraph Font','Hyperlink']:
    if name in word.styles:
        style=word.styles[name];style.font.name=family;style.font.size=Pt(11)
        style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
        if name.startswith('Heading'):style.font.bold=False;style.font.italic=False
for p in word.paragraphs:
    pf=p.paragraph_format;pf.line_spacing=Pt(LEAD);pf.space_before=Pt(0);pf.space_after=Pt(0)
    pf.first_line_indent=Pt(11.955);pf.keep_with_next=False;pf.page_break_before=False
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        r.font.name=family;r.font.size=Pt(11);r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    for t in p._p.xpath('.//w:t'):
        t.text=remap(t.text)
ps=word.paragraphs
ps[0].text=src_pars[0].text.upper()
ps[0].paragraph_format.space_before=Pt(18)
ps[0].paragraph_format.space_after=Pt(18)
for r in ps[0].runs:r.font.name=family;r.font.size=Pt(11);r.bold=True
ps[1].clear()
for text,sup in [('By Atharva Sheersh Pandey','1,a'),(', Adhyan Jain','1,b'),(',\nand Poornima Nedunchezhian','1,c')]:
    run=ps[1].add_run(text);run.font.small_caps=True;run.font.name=family;run.font.size=Pt(11)
    run=ps[1].add_run(sup);run.font.superscript=True;run.font.name=family;run.font.size=Pt(8)
ps[2].text='1 School of Computer Science and Engineering, VIT Vellore'
ps[3].text='a atharva.sheersh2024@vitstudent.ac.in; b adhyan.jain2024@vitstudent.ac.in'
ps[4].text='c Corresponding author: Poornima Nedunchezhian (poornima.n@vit.ac.in)'
ps[5]._p.getparent().remove(ps[5]._p)
for p in ps[:5]:
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.keep_with_next=True
for p in ps[2:5]:
    p.paragraph_format.line_spacing=Pt(9.9626);p.paragraph_format.space_after=Pt(9)
    for r in p.runs:r.font.name=family;r.font.size=Pt(8);r.italic=True
ps[6].text=src_pars[6].text.removeprefix('Abstract: ')
ps[6].paragraph_format.left_indent=Pt(64.765);ps[6].paragraph_format.right_indent=Pt(64.765)
ps[6].paragraph_format.line_spacing=Pt(SLEAD);ps[6].paragraph_format.space_after=Pt(19.925)
ps[6].paragraph_format.keep_with_next=False
for r in ps[6].runs:r.font.name=family;r.font.size=Pt(9)
# Keywords become a true first-page footnote, matching the reference pattern.
kwparagraph=ps[7];kwparagraph._p.getparent().remove(kwparagraph._p)

main=sub=0
for p in list(word.paragraphs):
    if p.style.name.startswith('Heading'):
        text=p.text
        if text=='REFERENCES':
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(16)
            p.paragraph_format.space_after=Pt(8);p.paragraph_format.first_line_indent=Pt(0)
            continue
        if text=='DECLARATIONS':p._p.getparent().remove(p._p);continue
        if text.startswith('APPENDIX A.'):
            continue
        is_sub=p.style.name=='Heading 2'
        if is_sub:sub+=1;label=f'{main}.{sub}. {text}.'
        elif re.match(r'^\d+ ',text):main=int(text.split()[0]);sub=0;label=re.sub(r'^(\d+) ',r'\1. ',text)+'.'
        else:label=text.capitalize()+'.'
        nxt=p._p.getnext()
        p.clear();rr=p.add_run(label+'  ');rr.bold=not is_sub;rr.italic=is_sub;rr.font.name=family;rr.font.size=Pt(11)
        p.paragraph_format.space_before=Pt(10 if is_sub else 13)
        p.paragraph_format.keep_with_next=False
        if nxt is not None and nxt.tag==qn('w:p'):
            np=WordParagraph(nxt,word)
            if np.style.name=='Normal' and np.text.strip():
                for child in list(nxt):
                    if child.tag!=qn('w:pPr'):p._p.append(deepcopy(child))
                nxt.getparent().remove(nxt)
    elif p.style.name=='Caption':
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.line_spacing=Pt(SLEAD);p.paragraph_format.space_before=Pt(5)
        p.paragraph_format.space_after=Pt(12)
        for run in p.runs:run.font.size=Pt(9);run.italic=True
        m=re.match(r'Table (\d+)\s+(.*)',p.text)
        if m:
            p.text='Table '+m[1]+'\n'+m[2];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            for run in p.runs:run.font.name=family;run.font.size=Pt(9);run.italic=True
        else:
            p.text=re.sub(r'^Fig\. (\d+) ',r'Fig \1. ',p.text)
            for run in p.runs:run.font.name=family;run.font.size=Pt(9);run.italic=True
for p in word.paragraphs:
    if p._p.xpath('.//w:drawing'):
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(12)
        p.paragraph_format.line_spacing=1.0
        p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for shape in word.inline_shapes:
    ratio=shape.height/shape.width;shape.width=Pt(W);shape.height=Pt(W*ratio)
for ti,t in enumerate(word.tables,1):
    t.autofit=False
    widths={1:[91,144,169.5],2:[92,312.5],3:[145,62,105,92.5],4:[117,287.5],5:[28,55,55,194.5,72],6:[106,71,83,65,79.5],7:[34,75,75,220.5],8:[30,68,68,80,74,84.5],9:[54,66,78,65,76,65.5],10:[60,55,67,82,73,67.5],11:[86,57,67,62,73,59.5]}[ti]
    for col,width in zip(t.columns,widths):col.width=Pt(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'nil');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Pt(widths[ci]);pr=cell._tc.get_or_add_tcPr()
            for sh in list(pr.findall(qn('w:shd'))):pr.remove(sh)
            for old in list(pr.findall(qn('w:tcBorders'))):pr.remove(old)
            if ri in [0,len(t.rows)-1]:
                cb=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'000000');cb.append(b);pr.append(cb)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.line_spacing=Pt(9.5 if len(widths)>=5 or ti==1 else 10.9589)
                p.paragraph_format.keep_with_next=ri<len(t.rows)-1;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:run.font.name=family;run.font.size=Pt(8 if len(widths)>=5 or ti==1 else 9);run.bold=False
                for tx in p._p.xpath('.//w:t'):tx.text=remap(tx.text)

# Replace bibliography with the same alphabetical order and citation mapping.
for p in list(word.paragraphs):
    if re.match(r'^\[\d+\]',p.text):p._p.getparent().remove(p._p)
for old,text,author in refs:
    p=word.add_paragraph()
    p.paragraph_format.left_indent=Pt(36);p.paragraph_format.first_line_indent=Pt(-36)
    p.paragraph_format.line_spacing=Pt(SLEAD);p.paragraph_format.space_after=Pt(1)
    p.add_run(f'[{mapping[old]}]  ')
    ar=p.add_run(author);ar.font.small_caps=True
    yr=re.search(r'(?<!\d)(19|20)\d{2}(?!\d)',text);year=yr[0] if yr else None
    tail=text[len(author):].lstrip(', ')
    if year:tail=re.sub(r'(?<!\d)'+year+r'(?!\d)[,.]?\s*','',tail,count=1)
    tail=tail.replace('“','').replace('”','').replace('"','')
    p.add_run((f' ({year}). ' if year else '. ')+tail)
    for run in p.runs:run.font.name=family;run.font.size=Pt(9)
# The supplied reference puts appendices after references.
pars=list(word.paragraphs)
ai=next(i for i,p in enumerate(pars) if p.text.startswith('APPENDIX A.'))
for p in pars[ai:ai+3]:
    word.element.body.remove(p._p);word.element.body.insert(len(word.element.body)-1,p._p)
ap=next(p for p in word.paragraphs if p.text.startswith('APPENDIX A.'))
ap.text='APPENDIX A: REPRODUCIBILITY';ap.alignment=WD_ALIGN_PARAGRAPH.CENTER
ap.paragraph_format.first_line_indent=Pt(0);ap.paragraph_format.space_before=Pt(16)
for run in ap.runs:run.font.name=family;run.font.size=Pt(11);run.bold=False

def page_field(p):
    r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
for sec in word.sections:
    for hdr,align in [(sec.header,WD_ALIGN_PARAGRAPH.RIGHT),(sec.even_page_header,WD_ALIGN_PARAGRAPH.LEFT)]:
        p=hdr.paragraphs[0];p.clear();p.alignment=align;p.paragraph_format.first_line_indent=Pt(0);page_field(p)
        for r in p.runs:r.font.name=family;r.font.size=Pt(10)
    for ftr in [sec.footer,sec.even_page_footer]:
        for p in ftr.paragraphs:p.clear()
    p=sec.first_page_footer.paragraphs[0];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;page_field(p)
    for r in p.runs:r.font.name=family;r.font.size=Pt(10)
word.core_properties.title=src_pars[0].text
word.core_properties.author=src_pars[1].text
docx=OUT/(STEM+'.docx');word.save(docx)

# Embed editable OpenType fonts in the Word package for portable typography.
from lxml import etree
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
with zipfile.ZipFile(docx) as z:parts={n:z.read(n) for n in z.namelist()}
ft=etree.fromstring(parts['word/fontTable.xml'])
fontentry=next((e for e in ft if e.get(qn('w:name'))==family),None)
if fontentry is None:fontentry=etree.SubElement(ft,qn('w:font'));fontentry.set(qn('w:name'),family)
relpath='word/_rels/fontTable.xml.rels'
rels=etree.fromstring(parts[relpath]) if relpath in parts else etree.Element('{http://schemas.openxmlformats.org/package/2006/relationships}Relationships')
for i,(style,tag) in enumerate([('Regular','embedRegular'),('Bold','embedBold'),('Italic','embedItalic'),('BoldItalic','embedBoldItalic')],1):
    key=uuid.uuid4();mask=key.bytes[::-1];raw=bytearray((TMP/f'TeXGyreTermesX-{style}.ttf').read_bytes())
    for j in range(32):raw[j]^=mask[j%16]
    filename=f'fonts/RevonTermesX{style}.odttf';parts['word/'+filename]=bytes(raw)
    rid=f'rIdRevonFont{i}'
    e=etree.SubElement(fontentry,qn('w:'+tag));e.set(qn('r:id'),rid);e.set(qn('w:fontKey'),'{'+str(key).upper()+'}')
    rel=etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship');rel.set('Id',rid);rel.set('Type',ns['r']+'/font');rel.set('Target',filename)
parts['word/fontTable.xml']=etree.tostring(ft,xml_declaration=True,encoding='UTF-8',standalone=True)
parts[relpath]=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=etree.fromstring(parts['[Content_Types].xml'])
if not any(e.get('Extension')=='odttf' for e in ct):
    e=etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default');e.set('Extension','odttf');e.set('ContentType','application/vnd.openxmlformats-officedocument.obfuscatedFont')
parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
# Real footnote on first page, with reference marker deliberately hidden.
foot=etree.Element(qn('w:footnotes'),nsmap={'w':ns['w']})
for fid,kind in [(-1,'separator'),(0,'continuationSeparator')]:
    f=etree.SubElement(foot,qn('w:footnote'));f.set(qn('w:type'),kind);f.set(qn('w:id'),str(fid))
    p=etree.SubElement(f,qn('w:p'));rr=etree.SubElement(p,qn('w:r'));etree.SubElement(rr,qn('w:'+kind))
f=etree.SubElement(foot,qn('w:footnote'));f.set(qn('w:id'),'1')
p=etree.SubElement(f,qn('w:p'));pp=etree.SubElement(p,qn('w:pPr'));sp=etree.SubElement(pp,qn('w:spacing'));sp.set(qn('w:line'),'219');sp.set(qn('w:lineRule'),'exact')
for txt,italic in [('Keywords and phrases: ',True),(kw+'.',False)]:
    rr=etree.SubElement(p,qn('w:r'));rp=etree.SubElement(rr,qn('w:rPr'));fonts=etree.SubElement(rp,qn('w:rFonts'));fonts.set(qn('w:ascii'),family);fonts.set(qn('w:hAnsi'),family)
    sz=etree.SubElement(rp,qn('w:sz'));sz.set(qn('w:val'),'18')
    if italic:etree.SubElement(rp,qn('w:i'))
    tx=etree.SubElement(rr,qn('w:t'));tx.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tx.text=txt
parts['word/footnotes.xml']=etree.tostring(foot,xml_declaration=True,encoding='UTF-8',standalone=True)
dr=etree.fromstring(parts['word/_rels/document.xml.rels'])
e=etree.SubElement(dr,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship');e.set('Id','rIdRevonFootnotes');e.set('Type',ns['r']+'/footnotes');e.set('Target','footnotes.xml')
parts['word/_rels/document.xml.rels']=etree.tostring(dr,xml_declaration=True,encoding='UTF-8',standalone=True)
e=etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override');e.set('PartName','/word/footnotes.xml');e.set('ContentType','application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
tree=etree.fromstring(parts['word/document.xml'])
abspar=next(p for p in tree.findall('.//'+qn('w:body')+'/'+qn('w:p')) if ''.join(p.itertext()).startswith('Revon is a Python prototype'))
rr=etree.SubElement(abspar,qn('w:r'));rp=etree.SubElement(rr,qn('w:rPr'));etree.SubElement(rp,qn('w:vanish'))
fn=etree.SubElement(rr,qn('w:footnoteReference'));fn.set(qn('w:id'),'1')
parts['word/document.xml']=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(docx,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in parts.items():z.writestr(n,b)

assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
audit={'source_sha256':source_hash,'reference_sha256':hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),
       'tables':11,'figures':7,'references':31,'citation_mapping':mapping,
       'font_family':'TeXGyreTermesX','pdf_body_pt':FS,'pdf_leading_pt':LEAD,
       'page_pt':[612,792],'margins_pt':[L,R],'text_width_pt':W,
       'abstract_width_pt':274.97,'source_unchanged':True,
       'note':'PDF uses exact full Type 1 font programs and measured sizes. Word rounds font sizes to half points; pagination can differ. No arXiv identifier or ORCID fabricated.'}
(TMP/'audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
print('Created',pdf,docx)
