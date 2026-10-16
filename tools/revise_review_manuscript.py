"""Build the semantic review revision from the preserved manuscript baseline.

Requires a completed, audited new campaign and its data-derived analysis.
Typesetting is separate; no measurements or historical evidence are changed.
"""
from pathlib import Path
from hashlib import sha256
import json,re,csv,statistics
from docx import Document
from docx.shared import Inches

ROOT=Path(__file__).resolve().parents[1]
BUNDLE=ROOT/'evidence/review-revision-20261016'
CAMPAIGN=BUNDLE/'campaign'
OUT=ROOT/'output/review_revision'
FIGURES=ROOT/'paper/review-assets/figures'
ARCHIVE='Revon_Review_Revision_Reproducibility_20261016.zip'

def load(path):return json.loads(path.read_text(encoding='utf8'))
def number(value):return f'{value:.3f}' if abs(value)<100 else f'{value:,.2f}'

def build():
    findings=load(CAMPAIGN/'analysis.json')
    if not findings['audit']['passed']:raise ValueError('complete audited evidence required')
    doc=Document(ROOT/'paper/previous versions/review-baseline-20261016/Revon_Manuscript_Source.docx')
    base=list(doc.paragraphs)
    assert len(base)==175 and len(doc.tables)==11 and len(doc.inline_shapes)==7
    old_tables={i+1:[[c.text for c in row.cells] for row in t.rows] for i,t in enumerate(doc.tables)}
    old_table_nodes={i+1:t._tbl for i,t in enumerate(doc.tables)}
    old_widths={1:[91,144,169.5],2:[92,312.5],3:[145,62,105,92.5],4:[117,287.5],5:[28,55,55,194.5,72],6:[106,71,83,65,79.5],7:[34,75,75,220.5],8:[30,68,68,80,74,84.5],9:[54,66,78,65,76,65.5],10:[60,55,67,82,73,67.5],11:[86,57,67,62,73,59.5]}
    # Resolve old and new table numbers from actual document order only once.
    for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
        if p.text:p.text=re.sub(r'\bTable (\d+)\b',lambda m:f'@@OLDTABLE{m[1]}@@',p.text)
    widths={f'@@OLDTABLE{i}@@':w for i,w in old_widths.items()}

    def para(anchor,text,style='Normal'):
        p=doc.add_paragraph(text,style=style);anchor._p.addprevious(p._p);return p
    def heading(anchor,text):return para(anchor,text,'Heading 2')
    def table(anchor,key,caption,rows,col_widths):
        token=f'@@NEWTABLE_{key}@@';widths[token]=col_widths
        para(anchor,token+' '+caption,'Caption')
        t=doc.add_table(rows=len(rows),cols=len(rows[0]));anchor._p.addprevious(t._tbl)
        for row,values in zip(t.rows,rows):
            for cell,value in zip(row.cells,values):cell.text=str(value)
        return token
    def figure(anchor,n,caption):
        p=para(anchor,'');p.add_run().add_picture(str(FIGURES/f'Fig{n}.png'),width=Inches(5.62))
        para(anchor,f'Fig. {n} {caption}','Caption')

    v=findings['validation'];threshold=findings['threshold']
    base[0].text='Revon: Log and Merkle Differencing for Durable Structured Data Versioning'
    base[6].text=(
      'Abstract: Revon is a Python prototype for durable, linear-history key/value versioning. We study the choice between log aggregation and hash-pruned comparison over the same SQLite-backed fixed-depth Merkle trie. The contribution is a controlled evaluation of this choice and its workflow costs, rather than a new index or general-purpose database. Our author-run Linux evaluation of the JSON-repaired implementation comprises 516 executions and a separate 72-execution geometry/telemetry supplement. Eight of nine paired hybrid/Merkle timing intervals favour hybrid, while the dense case includes parity. A frozen threshold nevertheless selects a slower path in a separate control. The Windows evaluation comprises 93 cases, each run in a fresh worker process, with disjoint calibration and validation histories, controlled sensitivity experiments, Wine Quality and CO2 public sources, and repeated runs at 100,000 and 1,000,000 records. Its largest threshold candidate wins calibration and uses log throughout the tested intervals; this does not establish a benefit from switching or an optimal crossover. Published CO2 revisions provide observed record changes, while wine corrections remain generated. We report algorithm behaviour, CLI execution phases and whole-worker resources separately. All measured outputs and historical states are checked against semantic oracles. The results describe bounded versioning workflows on the recorded machines; they establish neither storage-engine superiority nor independent-party reproduction.')
    base[10].text=(
      'We investigate how two correct diff paths behave when they share the same fixed-depth Merkle hash trie, addressed operation history and durable representation. Path copying, cryptographic content addresses and SQLite transactions are established mechanisms. Our contribution is their explicit integration into an auditable comparison of selection policies, with matched semantic outputs and documented measurement boundaries. The scope is limited to the measured versioning operations and workloads, without implying lower costs for structured-data management generally.')
    base[11].text=(
      'Dolt combines SQL with a commit graph and Prolly Trees [2], [3]. ForkBase contributes POS-Tree/SIRI storage and fork semantics [15], while Noms describes history-independent chunked collections [14]. Revon offers a narrower linear-history key/value prototype. We do not claim feature parity, a new family of persistent indexes, or a generally faster storage engine. The current comparison uses Revon\'s in-process API and Dolt\'s CLI; a separate phase breakdown makes the remaining interface costs explicit.')
    base[13].text='RQ1: What latency, storage and resource trade-offs are observed for complete versioning workflows, and how do API/CLI boundaries limit cross-system interpretation?'
    base[14].text='RQ2: How well does an operation-count selector carry from calibration to unseen generated histories when compared with both forced log and forced Merkle on the same representation?'
    base[15].text='RQ3: How do controlled changes in workload structure, payload and fixed-trie geometry affect differencing costs on the measured histories?'
    base[17].text='Implementation contribution: Revon integrates established content addressing, path copying and atomic SQLite persistence while exposing both diff paths under one encoded-value contract. This integration is an engineering contribution, not a claim to invent these mechanisms.'
    base[18].text='Policy contribution: we specify and test an interpretable operation-count rule against its forced-path controls. The calibration choice is frozen before held-out evaluation; adverse choices and the possibility that always-log suffices are retained.'
    base[19].text='Evidence contribution: observed diff outputs and historical states are checked against separately constructed oracles. Executed sources, raw measurements, failures, environment records and data preparation are retained, including the JSON repair and telemetry limitations.'
    base[20].text='Evaluation contribution: the author-run Linux campaigns are analysed separately from the 93-case Windows evaluation with disjoint seeds, matched sensitivity controls, Wine Quality and CO2 public sources and repeated scale observations. No timing rows are pooled across these campaigns.'
    base[21].text='Boundary of the contribution: the experiments concern single-writer linear histories and bounded workloads. They do not establish an optimal selector, a general geometry optimum, query-engine superiority or reproduction by an independent team.'
    base[35].text=base[35].text.replace('The public-record extension uses generated updates to one release, and does not establish performance across diverse datasets, independent hardware, concurrent workloads or production systems.','The retained Linux public-record extension uses generated updates to one TLC release. The Windows evaluation uses generated wine corrections and an observed CO2 publication sequence. These bounded sources do not establish performance across diverse domains, independent hardware, concurrent workloads or production systems.')
    base[106].text=base[106].text.replace('At a million rows,','In the retained Linux million-row diagnostic,').replace('repeated million-row trials are still needed.','the repeated Windows scale runs are reported separately in Section 6.4.')

    # Expanded related work supplements, rather than deletes, the historical table.
    heading(base[36],'Capabilities and comparison boundaries')
    para(base[36],'@@NEWTABLE_related@@ compares system abstractions and the evidence actually used here. A documented capability is not a benchmark result, and a system not evaluated here is not thereby unsupported. The cited ForkBase and Noms APIs describe typed collections; this study does not evaluate a relational layer over them. Git4Data is treated as a preprint about a different abstraction level [29].')
    table(base[36],'related','System capabilities and evidence scope; documented features are not performance results',[
      ['System','Representation / history','Diff / persistence','Branches / queries','Reproducibility and use here'],
      ['Revon','Fixed-depth hash trie; linear history','Log or Merkle; SQLite objects','No named branches/merge; key/value API','Executed snapshots, semantic proofs and raw trials; measured here'],
      ['Dolt [2,3]','Prolly Trees; commit graph','Native version diff; persistent block store','Branches/merge; SQL','Official docs/source; CLI workflow measured, engine not isolated'],
      ['ForkBase [15]','POS-Tree/SIRI; forkable versions','Hash-pruned diff/merge; addressed chunks','Fork on demand/conflict; typed KV APIs','Primary paper and APIs; no experiment run here'],
      ['Noms [14]','History-independent chunks; Prolly collections','Hash-based diff/sync/merge; chunk store','Typed collection API; not benchmarked here','Technical overview/source; conceptual comparison'],
      ['Git4Data [29]','MatrixOne immutable objects/MVCC; table versions','Database-native diff/merge','Snapshot/tag/branch; SQL extensions','Preprint/BranchBench claims belong to its authors; not run here']
    ],[53,91,91,78,91.5])

    # Keep the main methods readable; exact historical details move to Appendix B.
    moved={i:base[i].text for i in [64,69,74,75,76,77,78,80,81,82]}
    base[62].text='The retained primary campaign ran on a separate Ubuntu n2-standard-4 cloud VM; @@OLDTABLE2@@ identifies its environment and two source snapshots. The Windows campaign ran on a separate author-operated host. Differences in hardware, OS, runtime and storage remain confounded, and neither campaign is an independent-party reproduction.'
    base[64].text='Each campaign retains source hashes, a frozen protocol, environment records, correctness proofs and raw observations. The Linux preflight gates passed 67 and 69 tests; the Windows preflight gate passed 77. Packaging failures and the public-data premeasurement exclusion remain documented. Detailed records and commands are described in Appendices A and B.'
    base[74].text='The workflow is preparation, repository construction, timed operations, semantic verification and analysis. Initial import, commit, diff and checkout have separate boundaries. Diff outputs use the same sorted key and old/new-value-hash contract. @@NEWTABLE_timing@@ separates algorithm work from interface and worker costs; correctness checks and decision probes are outside operation timers.'
    base[75].text='For each execution, we compare the observed endpoint diff and every historical state with independent workload oracles. Revon also checks reverse/identity requests, stored-object integrity and reopen behaviour. Audits reconstruct the plan and semantic hashes; they verify internal consistency, not a fresh timing experiment.'
    base[76].text='The historical Linux estimates use seven fresh-repository trials and paired percentile-bootstrap timing intervals for fixed histories. The Windows campaign records seven cached queries within each fresh repository and reports independent-history or repository medians separately. These experimental units are not interchangeable; neither estimates a population of machines.'
    base[77].text='Resource sampling covers the worker and observed descendants, including setup and verification. It is not product-only memory or operation-only CPU. The primary Linux read-counter ambiguity is retained; the Linux supplement and Windows sampler distinguish supported zero counters from missing counters. Windows process I/O bytes are not direct physical-disk traffic.'
    base[78].text='Repository footprint measures persistent files, including metadata and representation costs. The Linux measurement is taken before closing/compacting; the Windows measurement follows close and the specified verification steps. These differently conditioned values are not pooled or interpreted as pure tree overhead.'
    base[80].text='An existing Windows calibration/held-out study selected the largest candidate, T=16,384; it did not identify a crossover. The Linux threshold controls preserve that choice. The Windows campaign uses a separately frozen design, source snapshot and disjoint seeds, described below; its threshold results are analysed separately.'
    base[81].text='The retained Linux controls compare forced log, forced Merkle and T=16,384 hybrid at 8,192 and 32,768 operations over 10,000 rows. The Windows measurements do not retune these controls.'
    base[82].text='The Linux geometry supplement evaluates five configurations on one fixed spread workload. The local sensitivity experiment adds three generated histories and matched equal-capacity geometries. Bucket capacity and operation threshold remain distinct parameters; neither study establishes a universal geometry optimum.'
    table(base[75],'timing','Timing and interpretation boundaries',[
      ['Quantity','Revon / same-representation control','Dolt CLI workflow','Interpretation'],
      ['Diff latency','Version API, selected path and canonical output','Native JSON-diff command, parsing and canonical output','Complete interface contract; no engine-only ranking'],
      ['Phase probe','Separate ancestry/count probe; includes detached-view copying','Command, JSON parse and output-construction timers','Neither probe nor command isolates pure tree-engine work'],
      ['Correctness','All states, reverse/identity, integrity and reopen outside timers','Observed native rows and historical checkout outside timers','Observed values compared with an independently constructed oracle'],
      ['Resources','Worker plus sampled children, including oracle/setup','Worker plus sampled Dolt processes','10 ms samples; short processes/peaks may be missed'],
      ['Persistent bytes','Closed local files; old Linux conditioning retained','Repository directory files','Includes metadata/indexes; not isolated algorithm overhead']
    ],[61,113,113,117.5])
    heading(base[83],'Separate local calibration and validation')
    para(base[83],'The local campaign fixes N=12,000, H=8, payload=32 bytes and b=8,d=4 for the primary selector comparison. Two calibration seeds and three different validation seeds cross spread/repeated-key histories with 64, 512 and 4,096 changes per commit. Forced log, forced Merkle and T=512/4,096/16,384/65,536 hybrid share a durable representation within each case. Two warm-ups precede seven measured query blocks; configuration order is deterministically shuffled. Each case starts in a fresh worker, while queries within it share resident objects and filesystem cache. No cold-cache condition is claimed.')
    para(base[83],'Calibration minimises the equally weighted mean of hybrid case medians normalised by the lower forced-path median; exact ties choose the smaller T. The selected parameter and UTC time are written before the first validation worker starts. Validation outcomes do not select a new winner. Sensitivity changes rows, history interval, payload, geometry and routing in a prespecified staged design. At fixed accumulated operations, interval length and changes per commit trade off; the design is not a full factorial. Appendices and retained protocols give all seeds and matched controls.')
    table(base[83],'protocol','Windows experiment matrix; cached queries are distinct from independent history or repository repetitions',[
      ['Component','Fresh workers','Histories / control','Warm-up queries','Measured queries'],
      ['Calibration','12','2 seeds × 2 localities × 3 densities','144','504'],
      ['Held-out','18','3 different seeds × same profiles','216','756'],
      ['Sensitivity','21','3 seeds × 7 matched variations','252','882'],
      ['Monitor diagnostic','6','3 off/on pairs, one fixed history','72','252'],
      ['Public sources','18','2 sources × 3 runs × M/H/Dolt','36','126'],
      ['Scale','18','2 sizes × 3 seeds × M/H/Dolt','36','126'],
      ['Total','93','Different campaigns are never pooled','756','2646']
    ],[72,50,160,61,61.5])
    heading(base[83],'Additional public histories and repeated scale')
    para(base[83],'Wine Quality contributes 6,497 red/white records, keyed by colour and source row ordinal [32]. Eight batches of 64 seeded quality-field corrections are generated for each of three histories; they are not observed specimen revisions. The CO2 package supplies published NOAA-derived CSV revisions [33]. We retain the eleven most recent file-changing snapshots acquired before timing and exclude one header-only publication failure before freezing the experiment. Ten populated snapshots give nine observed transitions, initially 812 records. The source contains unnamed trailing columns, which are preserved as additional fields. Dates identify records; CSV row order and formatting are discarded. The result is a published repository history, not NOAA transaction logs or atmospheric-science evidence.')
    para(base[83],'Three CO2 repositories repeat the same observed history, whereas wine and scale use distinct generated seeds. Each source is measured through Revon-M, T=4,096 Revon-H and native Dolt JSON diff, rotating model order across runs. The scale study holds H=4, C=64 and payload=32 fixed at 100,000 and 1,000,000 records, with three seeds per size/system. Medians and ranges of repository medians remain descriptive. Source URLs, licences, versions, byte checksums and transformation rules accompany the data; a bounded publication history was chosen instead of bulk Software Heritage acquisition on this host.')
    para(base[83],'The complete count matrix is retained in @@NEWTABLE_protocol@@ in Appendix B; exact public-data checksums and the original Linux public protocol are in @@OLDTABLE4@@ there. The main text retains timing boundaries and experimental units needed to interpret the results.')

    # Local results are inserted after all seven preserved Linux figures.
    anchor=base[117];para(anchor,'6 Windows Evaluation Results','Heading 1')
    heading(anchor,'Held-out selector behaviour')
    para(anchor,f"Calibration selected T={threshold:,}, the largest candidate. Its normalised calibration score was {findings['calibration_scores'][str(threshold)]:.3f}. On 18 held-out histories, its mean normalised latency was {v['calibrated_normalized_mean']:.3f}, versus {v['old_normalized_mean']:.3f} for T=4,096. Agreement with the lower forced-path median was {v['calibrated_agreement']}/18 and {v['old_agreement']}/18 respectively. These are descriptive choices relative to noisy forced-path medians, not classification accuracy against a known optimal strategy. @@NEWTABLE_selector@@ and Figure 8 report all profile groups rather than selecting only favourable cases.")
    para(anchor,'All tested intervals contain at most 32,768 operations, so T=65,536 requests log throughout this matrix. Thus the calibration winner is behaviourally an always-log policy over these intervals. Its selection cannot establish that adaptive switching is beneficial or locate a global crossover. Hybrid/forced ratios below one may reflect timing variation despite equivalent selected paths. The earlier adverse Linux result is retained unchanged. Production defaults are not altered by this experimental selection.')
    sel=list(csv.DictReader((CAMPAIGN/'selector.csv').open(encoding='utf8')))
    selector_rows=[['Locality','Operations','Forced log ms','Forced Merkle ms','H4096 ms','Calibrated ms']]
    for loc in ['spread','repeated-key']:
        for operations in [512,4096,32768]:
            rows=[r for r in sel if r['suite']=='validation' and r['locality']==loc and int(r['operations'])==operations]
            selector_rows.append([loc,f'{operations:,}']+[number(statistics.median(float(r[k]) for r in rows)) for k in ['log_ms','merkle_ms','hybrid4096_ms','calibrated_ms']])
    table(anchor,'selector','Held-out median of three history medians; fixed N/H/payload and no population interval',selector_rows,[68,60,69,69,69,69.5])
    figure(anchor,8,'Held-out selector comparison at N=12,000 and H=8. Lines show medians across three generated-history medians; shaded bands span their minimum and maximum. Each history has seven cached measured queries. Bands are descriptive ranges, not 95% confidence intervals. The calibrated threshold uses log at every shown operation count. Forced paths share the same durable representation.')
    heading(anchor,'Sensitivity and resource interpretation')
    cost=json.loads((ROOT/'evidence/selector-cost-20261016/summary.json').read_text())
    para(anchor,f"The recorded decision probe is outside diff timing and reads a detached public changeset view, so its copying work is not pure selector overhead. A separate post-campaign diagnostic extracts the unchanged production function prefix through input validation, ancestry and internal operation counting, stopping before diff dispatch. Across the 18 held-out histories and two thresholds, 252,000 measured warm prefix calls give a median of case/threshold medians of {cost['case_threshold_median_us']:.3f} microseconds (range {cost['min_case_threshold_median_us']:.3f}-{cost['max_case_threshold_median_us']:.3f}). The extracted decision agrees with production dispatch. This microbenchmark excludes diff traversal and detached-view copying; it does not isolate cold-cache overhead or provide a causal amount to subtract. The original 93-case records are unchanged.")
    para(anchor,'The 21 sensitivity cases change one prespecified factor or matched control across three seeds. Payload=256 bytes, N=24,000, fixed-operation H=16, equal-capacity b4-d6/b16-d3 geometries, and a matched C=8 hash-route/spread pair are all retained in the selector summary. Geometry cases match the base history state hashes; only routing geometry changes. The fixed-operation interval comparison trades H against changes/commit and does not isolate both independently. These results concern the tested profiles and do not select a generally best geometry.')
    sensitivity=[r for r in sel if r['suite']=='sensitivity']
    snippets=[]
    for suffix in ['rows24000','history16','payload256','geometry4x6','geometry16x3','route8','route-control8']:
        values=[float(r['merkle_ms']) for r in sensitivity if r['case'].endswith('-'+suffix)]
        snippets.append(f'{suffix}: {statistics.median(values):.3f} ms')
    para(anchor,'Forced-Merkle medians across the three history medians were '+('; '.join(snippets))+'. Full per-history timing, chosen paths and resource scopes remain available in the retained summaries. These exploratory comparisons receive no multiple-testing significance claim.')
    monitor=findings['monitor_diagnostic'];ratios=[r['on_over_off'] for r in monitor if r['strategy']=='hybrid' and r['threshold']==4096]
    para(anchor,f"The three sampler off/on pairs gave T=4,096 median-latency ratios ranging from {min(ratios):.3f} to {max(ratios):.3f}. Off precedes on within each pair, and background/cache conditions are uncontrolled; this diagnostic does not identify a causal overhead correction and is not subtracted from other timings. Resource counters span whole workers. Tests reject the replacement of observed zero I/O by missing data, but Windows counters still cannot identify database logical reads or direct physical-disk traffic.")
    heading(anchor,'Public data and interface phases')
    public_rows=[['Source / system','Runs','Diff ms','Range ms','Import s','Persistent MiB']]
    for r in findings['public_and_scale']:
        if r['dataset'] not in ['co2-observed','wine-generated']:continue
        public_rows.append([r['dataset'].split('-')[0]+'/'+r['model'],r['repositories'],number(r['diff_median_ms']),f"{number(r['diff_min_ms'])}-{number(r['diff_max_ms'])}",number(r['import_median_s']),number(r['storage_median_mib'])])
    table(anchor,'public','Additional public sources: median and range of three repository medians; CO2 repeats one history',public_rows,[91,49,65,70,60,69.5])
    figure(anchor,9,'Additional public data. Points show medians of three repository medians and bars show their full range. CO2 repeats the same observed publication history; wine uses three generated correction histories. Each repository contributes seven cached queries. The API/CLI interfaces differ; these are workflow latencies, not isolated storage-engine measurements.')
    for r in findings['public_and_scale']:
        if r['dataset'] in ['co2-observed','wine-generated'] and r['model']=='dolt':
            para(anchor,f"For {r['dataset']}, the Dolt command phase had a median of {r['command_ms']:.3f} ms, JSON interpretation {r['parse_ms']:.3f} ms and common output construction {r['canonical_ms']:.3f} ms. The command accounts for {100*r['command_fraction']:.1f}% of the median complete diff latency, while combining process startup, native work and native JSON emission. The decomposition cannot assign that whole interval to a storage engine, nor justify subtracting a separately measured startup constant.")
    heading(anchor,'Repeated local scale observations')
    scale_rows=[['Rows / system','Import s','Commit ms','Diff ms','Checkout s','Worker RSS GiB']]
    for r in findings['public_and_scale']:
        if r['dataset'] not in ['100000','1000000']:continue
        scale_rows.append([f"{int(r['dataset']):,}/{r['model']}"]+[number(r[k]) for k in ['import_median_s','commit_median_ms','diff_median_ms','checkout_median_s','rss_median_gib']])
    table(anchor,'scale','Repeated Windows scale study: medians of three seed/repository summaries; full ranges in evidence',scale_rows,[100,60,64,60,60,60.5])
    figure(anchor,10,'Local scale observations with three generated histories and fresh repositories per size/system. Curves connect descriptive medians at two sizes, not fitted asymptotic scaling laws. Diff queries are cached; import and checkout have separate boundaries. Sampled peak RSS includes workload/oracle generation, storage work, verification and sampled children. The Linux million-row n=1 diagnostics remain a different campaign.')
    para(anchor,'The repeated Windows scale runs are analysed separately from the Linux diagnostics. The three generated histories share the same initial synthetic distribution and differ in updates; they are not three independent public datasets. Memory availability and background processes are recorded in the manifest, not controlled away. These observations support conditional descriptions of the measured workflows, not a population-level ranking or a general scaling law.')
    heading(anchor,'Key takeaways')
    para(anchor,'RQ1: the workflow measurements expose latency, persistence and worker-resource trade-offs, with API/CLI costs still preventing an engine-only ranking. RQ2: some fixed histories benefit from log relative to forced Merkle, but the Windows calibration winner uses log everywhere tested; optimal switching remains unestablished. RQ3: controlled sensitivity histories make payload, routing and geometry effects visible within the tested design, without identifying a universally best layout. Observed CO2 transitions broaden revision realism, while generated wine and synthetic histories retain their own scope limits.')

    base[117].text='7 Discussion';base[122].text='8 Limitations And Future Work';base[129].text='9 Conclusion'
    base[118].text='RQ1 concerns the complete versioning contract. The Linux and local experiments both retain import, commit, diff, checkout and persistent footprint, but their rows are analysed separately. Snapshot and Log-only expose representation costs; the M/H/log controls hold the Revon representation fixed. Dolt provides a richer SQL/versioning interface. Its command-phase decomposition improves transparency without making its native tree cost comparable to an isolated Python tree call.'
    base[119].text='RQ2 concerns transfer to unseen histories, not merely repeated timing of one history. The Windows split controls primary N, H, payload and geometry and freezes selection before validation. Its largest candidate behaves as always-log within the measured interval range. The adverse Linux selection and dense-case parity remain useful counterexamples. A more complex selector would require new calibration and untouched testing; these results do not justify adding complexity or changing the production threshold.'
    base[120].text='RQ3 concerns controlled changes to the fixed trie and its workload. Equal-capacity geometries can differ in traversal depth and ancestor reconstruction while having equal expected occupancy under uniform routing. Payload encoding, touched leaves and output materialisation also contribute. The Windows sensitivity set uses independent seeds and matched controls alongside the Linux study of one fixed history, but does not estimate every factor interaction or a population optimum.'
    base[123].text='External validity: Linux results come from one author-run cloud VM, and Windows results from one author-run host. Both campaigns were conducted by the authors; the Windows campaign does not reproduce the Linux environment. Hardware, OS, runtime, virtualization and storage vary together between campaigns. Background services and filesystem caching remain active; the local manifest records limited available RAM and process conditions. No entirely idle, cold-cache or OS-causal claim follows.'
    base[124].text='Workload validity: the original TLC histories remain generated from one monthly release, with n=1 Linux scale diagnostics. The Windows evaluation includes three held-out generated seeds per primary profile, generated wine corrections and one observed CO2 publication sequence. Three CO2 repositories repeat that sequence rather than independently sampling its history. A header-only upstream revision is excluded before timing, and unnamed columns are retained without scientific interpretation. Repeated synthetic scale histories broaden observations but share their initial distribution. More domains, observed histories and hardware are still needed for broad generalisation.'
    base[126].text='Measurement and inference: process startup, native work and JSON emission remain combined in Dolt\'s command interval. Whole-worker RSS includes oracle/history state and is not product-only memory. Ten-millisecond sampling can miss short-lived processes, and Windows I/O counters do not identify physical-disk or database logical reads. The off/on sampler diagnostic also confounds execution order and cache state. The historical bootstrap intervals concern fixed-history timings; Windows three-history ranges are descriptive and have no multiple-comparison correction. Source and semantic-hash checks establish internal consistency, not independent performance reproduction.'
    base[127].text='Selection and design: calibration is limited to four thresholds, two generated seeds and the declared primary profiles. T=65,536 is the largest candidate and does not switch within the tested operation range. That outcome leaves a broadly useful adaptive selector unestablished. The staged sensitivity design is not a full factorial; H and density trade off at fixed operation count, and hash-route tests require their own low-density control. A new selector or crossover claim requires a separately frozen experiment.'
    base[130].text=(
      f'Revon integrates established fixed-trie content addressing, changesets and atomic SQLite persistence to expose two semantically consistent ways to compute a version diff. The retained Linux evidence comprises 516 executions and a 72-execution supplement, including dense-case uncertainty and an adverse threshold choice. The Windows evaluation comprises 93 worker cases and 3,402 checked queries, with held-out generated histories, Wine Quality and CO2 public sources, interface-phase measurements and repeated scale observations. Calibration selects T={threshold:,}, which behaves as always-log in this matrix and does not establish optimal switching. The combined evaluation covers controlled workload variations and public data while retaining API/CLI, sampling, dataset and hardware limits. The contribution is an auditable prototype and a bounded empirical trade-off study; a new general-purpose database, universal speed advantage and independent-party reproduction are not established.')
    base[132].text=f'Readers can find the code and experiment records in the Revon GitHub repository. The paper/ directory contains the original Linux reproducibility archive and {ARCHIVE}, which retains the Windows campaign, source snapshots, raw observations, protocols and internal audits. The README files explain dependency setup, public-source acquisition, execution and verification [24]. Repository updates do not replace the frozen source inside an experiment bundle.'
    base[133].text='The two Linux bundles and the Windows study belong to different campaigns and must not be pooled. Original TLC records and generated repositories are rebuilt using the recorded preparation steps. The local public-source manifest records licences, acquisition versions, checksums and the premeasurement CO2 exclusion. Auditors run on disposable copies with localised data paths, preserving collected bytes. Internal audits do not independently reproduce timing; matching semantic hashes do not guarantee identical elapsed times.'
    base[140].text=f'Data and code availability: NYC TLC provides the original public taxi source [31]; UCI Wine Quality [32] and the published CO2 package [33] provide the additional sources. Code, frozen experiment snapshots and evidence are available at https://github.com/atharvasheersh/Revon. The Windows reproducibility archive is {ARCHIVE}; the original Linux archive remains available separately. Original TLC records and generated databases are not redistributed. The public-source manifest records licences and transformations; Appendix A describes access.'
    # Detailed methods stay after availability and before declarations in source;
    # the renderer moves all appendices together after the bibliography.
    para(base[134],'APPENDIX B. METHODS AND EVIDENCE DETAILS','Heading 1')
    for i in [64,74,75,76,77,78,80,81,82]:para(base[134],moved[i])
    para(base[134],'For the local study, refer to the frozen protocol and full case list in review-revision-20261016/campaign. Calibration seeds 202610161/202610162 are distinct from validation/sensitivity seeds 202610171–202610173. Workers use PYTHONHASHSEED=20261016. The monitor condition uses seed 202610190; public generated corrections use seeds 202610201–202610203; scale uses seeds 202610221–202610223. Queries share a repository within each worker, and every new worker builds fresh persistent state. The chosen threshold is timestamped before validation. There is no post-validation selection or pooled Linux/Windows summary.')
    para(base[134],'The Windows Dolt instrumentation invokes the same native JSON-diff command and parser contract as the baseline adapter, with explicit phase timers. A preflight compares its actual output against the native adapter and the independent semantic oracle. The Windows Revon queries verify observed output hashes, all historical state hashes, reverse/identity requests, object integrity and reopen behaviour outside timing. Failed setup checks and data-quality exclusions precede measurement and remain documented. The exact query matrix and timing source hashes are retained; changed source requires a different campaign.')
    para(base[134],'A portable audit copies evidence to a disposable directory, imports the executed snapshot in a fresh interpreter, redirects only the recorded absolute public-data location, regenerates workloads and state/diff oracles, checks every configuration/trial/order identity and recomputes the numerical summaries. It then confirms that the original collection hashes remain intact. The archive provides these checks and acquisition recipes; they validate evidence consistency rather than elapsed times or a new research team\'s reproduction.')
    # Technical checksum and full count tables belong to the methods appendix.
    original_checksum=old_table_nodes[4]
    for node in [base[72]._p,original_checksum]:base[134]._p.addprevious(node)
    protocol_caption=next(p for p in doc.paragraphs if p.text.startswith('@@NEWTABLE_protocol@@'))
    protocol_table=protocol_caption._p.getnext()
    for node in [protocol_caption._p,protocol_table]:base[134]._p.addprevious(node)
    doc.add_paragraph('[32] P. Cortez, A. Cerdeira, F. Almeida, T. Matos and J. Reis, "Wine Quality [Dataset]," UCI Machine Learning Repository, 2009. https://doi.org/10.24432/C56S3T. CC BY 4.0.')
    doc.add_paragraph('[33] DataHub Contributors, "CO2 PPM - Trends in Atmospheric Carbon Dioxide," NOAA-derived published data package, accessed October 2026. https://github.com/datasets/co2-ppm. ODC-PDDL-1.0; acquired revision identifiers and file hashes are retained in the experiment manifest.')

    mapping={};table_index=0;last_token=None;ordered_widths={}
    for node in doc.element.body:
        if node.tag.endswith('}p'):
            text=''.join(node.xpath('.//w:t/text()'))
            m=re.match(r'(@@(?:OLDTABLE\d+|NEWTABLE_\w+)@@)',text)
            if m:last_token=m[1]
        elif node.tag.endswith('}tbl'):
            table_index+=1
            if not last_token:raise ValueError('table caption token missing')
            mapping[last_token]=table_index;ordered_widths[table_index]=widths[last_token];last_token=None
    assert table_index==17 and len(doc.inline_shapes)==10
    for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
        if p.text:
            text=p.text
            for token,n in mapping.items():text=text.replace(token,f'Table {n}')
            if '@@' in text:raise ValueError('unresolved table reference')
            if text!=p.text:p.text=text
    # The author requested removal of this declaration from the current paper.
    # Keep the historical baseline intact, but never restore it on rebuild.
    contributions=[p for p in doc.paragraphs if p.text.startswith('Author contributions:')]
    if len(contributions)!=1:raise ValueError('expected one author-contributions declaration in source')
    contributions[0]._p.getparent().remove(contributions[0]._p)
    OUT.mkdir(parents=True,exist_ok=True)
    source_path=OUT/'Revon_Review_Revised_Source.docx';doc.save(source_path)
    metadata={'source_sha256':sha256(source_path.read_bytes()).hexdigest(),'table_widths':ordered_widths,
      'historical_table_numbers':{i:mapping[f'@@OLDTABLE{i}@@'] for i in range(1,12)},
      'new_table_numbers':{k:mapping[f'@@NEWTABLE_{k}@@'] for k in ['related','timing','protocol','selector','public','scale']},
      'historical_tables':old_tables,'title':base[0].text,'figures':10,'tables':17,'references':33,'source':str(source_path.relative_to(ROOT))}
    (OUT/'publication_inputs.json').write_text(json.dumps(metadata,indent=2),encoding='utf8')
    print(json.dumps({'source':str(source_path),'tables':17,'figures':10,'references':33},indent=2))

if __name__=='__main__':build()
