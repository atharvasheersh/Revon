"""Failure-oriented controls for the expanded revision integrity gate."""
import copy,json,unittest,zipfile,shutil
import pymupdf
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from tools.check_review_campaign import ROOT,check_case_matrix
from tools.check_review_package import check as check_package
from tools.check_springer_paper import check_springer_doc,check_springer
from tools.check_final_paper import check_package as old_package

class ReviewRevisionChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        folder=ROOT/'evidence/review-revision-20261016/campaign'
        cls.protocol=json.loads((folder/'protocol.json').read_text());cls.case=cls.protocol['cases'][0]
        cls.rows=json.loads((folder/'cases'/cls.case['name']/'result.json').read_text())['records']

    def test_retained_matrix_passes(self):check_case_matrix(self.case,self.rows,self.protocol)

    def test_duplicate_configuration_is_rejected(self):
        rows=copy.deepcopy(self.rows);rows[1]=copy.deepcopy(rows[0])
        with self.assertRaisesRegex(ValueError,'matrix mismatch'):check_case_matrix(self.case,rows,self.protocol)

    def test_wrong_execution_order_is_rejected(self):
        rows=copy.deepcopy(self.rows);rows[0]['execution_order']=rows[1]['execution_order']
        with self.assertRaisesRegex(ValueError,'execution-order'):check_case_matrix(self.case,rows,self.protocol)

    def test_wrong_public_strategy_is_rejected(self):
        case={'suite':'public','model':'revon-m'}
        rows=[{'strategy':'hybrid','threshold':4096,'trial_kind':kind,'trial':i,'execution_order':1} for kind,n in [('warmup',2),('measured',7)] for i in range(1,n+1)]
        with self.assertRaisesRegex(ValueError,'matrix mismatch'):check_case_matrix(case,rows,self.protocol)

    def test_changed_new_result_is_rejected(self):
        doc=Document(ROOT/'paper/Revon Research Paper Springer.docx')
        spec=json.loads((ROOT/'paper/Revon_Review_Publication_Model_20261016.json').read_text())
        number=spec['new_table_numbers']['public'];doc.tables[number-1].cell(1,2).text='999.999'
        model,_,_=old_package(ROOT/'paper/Revon_Linux_Reproducibility_Supplement.zip',ROOT)
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'changed.docx';doc.save(path)
            with self.assertRaisesRegex(ValueError,f'Springer table {number} differs'):check_springer_doc(path,model)

    def test_removed_current_limitation_is_rejected(self):
        doc=Document(ROOT/'paper/Revon Research Paper Springer.docx')
        target=next(p for p in doc.paragraphs if p.text.startswith('Selection and design:'))
        target._p.getparent().remove(target._p)
        model,_,_=old_package(ROOT/'paper/Revon_Linux_Reproducibility_Supplement.zip',ROOT)
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'changed.docx';doc.save(path)
            with self.assertRaisesRegex(ValueError,'paragraph differs'):check_springer_doc(path,model)

    def test_missing_pdf_abstract_text_is_rejected(self):
        model,_,_=old_package(ROOT/'paper/Revon_Linux_Reproducibility_Supplement.zip',ROOT)
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'changed.docx';shutil.copy2(ROOT/'paper/Revon Research Paper Springer.docx',path)
            with pymupdf.open(ROOT/'paper/Revon Research Paper Springer.pdf') as pdf:
                rect=pdf[0].search_for('Revon is a Python prototype')[0]
                pdf[0].add_redact_annot(rect,fill=(1,1,1));pdf[0].apply_redactions();pdf.save(path.with_suffix('.pdf'))
            with self.assertRaisesRegex(ValueError,'PDF/DOCX paragraph differs'):check_springer(path,model)

    def test_new_package_tampering_is_rejected(self):
        with TemporaryDirectory() as tmp:
            path=Path(tmp)/'changed.zip'
            with zipfile.ZipFile(ROOT/'paper/Revon_Review_Revision_Reproducibility_20261016.zip') as source,zipfile.ZipFile(path,'w') as target:
                for item in source.infolist():target.writestr(item,b'changed' if item.filename=='evidence/review-revision-20261016/campaign/analysis.json' else source.read(item))
            with self.assertRaisesRegex(ValueError,'review package checksum mismatch'):check_package(path)
