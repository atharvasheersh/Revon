"""Editorial input guards must not conceal changed source or historical data."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import tempfile
import unittest
from docx import Document
from tools.first_submission_editorial import apply_editorial, historical_editorial_cells, HISTORICAL_CELL


class EditorialGuards(unittest.TestCase):
    def test_changed_source_is_rejected_before_any_mutation(self):
        doc=Document();doc.add_paragraph('first');doc.add_paragraph('unexpected')
        entries=[{'index':i,'source_text':s,'source_sha256':sha256(s.encode()).hexdigest(),'text':'replacement'} for i,s in enumerate(['first','second'])]
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'edits.json';path.write_text(json.dumps({'schema':1,'paragraphs':entries,'table_edits':[]}))
            with self.assertRaisesRegex(ValueError,'input changed'):apply_editorial(doc,path)
        self.assertEqual([p.text for p in doc.paragraphs],['first','unexpected'])

    def test_changed_table_is_rejected_before_any_mutation(self):
        doc=Document();doc.add_paragraph('first');doc.add_table(rows=1,cols=1).cell(0,0).text='unexpected'
        spec={'schema':1,'paragraphs':[{'index':0,'source_text':'first','source_sha256':sha256(b'first').hexdigest(),'text':'replacement'}],
              'table_edits':[{'table':1,'row':0,'column':0,'old':'expected','new':'replacement'}]}
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'edits.json';path.write_text(json.dumps(spec))
            with self.assertRaisesRegex(ValueError,'table input changed'):apply_editorial(doc,path)
        self.assertEqual(doc.paragraphs[0].text,'first')
        self.assertEqual(doc.tables[0].cell(0,0).text,'unexpected')

    def test_historical_exception_changes_only_the_exact_metadata_cell(self):
        table,row,col,old,new=HISTORICAL_CELL
        cells=[['label','value'] for _ in range(9)];cells[row][col]=old
        original=deepcopy(cells);expected=deepcopy(cells);expected[row][col]=new
        self.assertEqual(historical_editorial_cells(table,cells),expected)
        self.assertEqual(cells,original)
        self.assertEqual(historical_editorial_cells(table+1,cells),cells)
        altered=deepcopy(cells);altered[1][1]='changed numeric value 999'
        self.assertNotEqual(historical_editorial_cells(table,altered),expected)
        altered=deepcopy(cells);altered[row][col]=old+' changed'
        self.assertEqual(historical_editorial_cells(table,altered),altered)


if __name__=='__main__':unittest.main()
