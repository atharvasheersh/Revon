"""Negative controls for the current manuscript's numerical and link integrity."""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from docx import Document
from tools.check_final_paper import ROOT, check_package
from tools.check_springer_paper import check_springer, check_springer_doc, REPOSITORY


class SpringerPublicationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = ROOT / "paper/Revon Research Paper Springer Nature.docx"
        cls.model, _, _ = check_package(ROOT / "paper/Revon_Linux_Reproducibility_Supplement.zip", ROOT)

    def test_current_pair_passes(self):
        result=check_springer(self.source,self.model)
        self.assertEqual(result['references'],33 if (ROOT/'paper/Revon_Review_Publication_Model_20261016.json').exists() else 31)

    def mutate(self, edit, error):
        with TemporaryDirectory() as directory:
            doc = Document(self.source)
            edit(doc)
            path = Path(directory) / "modified.docx"
            doc.save(path)
            with self.assertRaisesRegex(ValueError, error):
                check_springer_doc(path, self.model)

    def test_modified_result_is_rejected(self):
        index=5 if (ROOT/'paper/Revon_Review_Publication_Model_20261016.json').exists() else 4
        self.mutate(lambda d: setattr(d.tables[index].cell(1, 1), "text", "999.999"), f"Springer table {index+1} differs")

    def test_missing_figure_is_rejected(self):
        def remove(d):
            node = d.inline_shapes[0]._inline
            node.getparent().remove(node)
        self.mutate(remove, "Springer figure count changed")

    def test_stale_commit_link_is_rejected(self):
        def pin(d):
            for rel in d.part.rels.values():
                if rel.target_ref == REPOSITORY:
                    rel._target = REPOSITORY + "/tree/" + "a" * 40
        self.mutate(pin, "Springer current GitHub links missing")
